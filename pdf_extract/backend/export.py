"""Excel / CSV export: per-template sheets, or one unified table via an output format."""
import csv
import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, i18n
from .extractor import to_number
from .i18n import bi

FILL_WARN = PatternFill("solid", fgColor="FFF2CC")
FILL_ERR = PatternFill("solid", fgColor="F8CBAD")
FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")

# Columns the export adds itself; they follow the UI language at export time.
COLS = {
    "batch": bi("批次号", "Batch No."),
    "file": bi("来源文件", "Source File"),
    "pages": bi("页码", "Pages"),
    "notes": bi("核对提示", "Review Notes"),
    "alloc": bi("分摊说明", "Allocation Note"),
}
# A unified-format column with one of these names is filled with the batch number unless mapped explicitly.
BATCH_COLUMNS = set(COLS["batch"].values())


def _col(key: str, lang: str) -> str:
    return i18n.text(COLS[key], lang)


# Each table is (title, columns, rows); a row is a list of (value, issue_level|None).

def _issue_levels(doc: dict) -> dict:
    levels = {}
    for i in doc.get("issues") or []:
        p = i["path"]
        if levels.get(p) != "error":
            levels[p] = i["level"]
    return levels


def _row_notes(doc: dict, row_idx: int | None, lang: str) -> str:
    msgs = []
    for i in doc.get("issues") or []:
        p = i["path"]
        if p.startswith("header.") or p == "items" or (row_idx is not None and p.startswith(f"items.{row_idx}.")):
            msgs.append(i18n.text(i["msg"], lang))
    return ("；" if lang == "zh" else "; ").join(dict.fromkeys(msgs))


def _resolve(source: str, doc: dict, tpl: dict, header: dict, row: dict | None, row_idx: int | None, levels: dict):
    if not source:
        return None, None
    if source == "meta.filename":
        return doc["filename"], None
    if source == "meta.template":
        return tpl["name"], None
    if source == "meta.pages":
        return ",".join(str(p) for p in doc.get("pages") or []), None
    if source == "meta.batch":
        return doc.get("batch_code", ""), None
    if source.startswith("const:"):
        return source[6:], None
    if source.startswith("header."):
        k = source[7:]
        return header.get(k), levels.get(f"header.{k}")
    if source.startswith("item."):
        k = source[5:]
        if row is None:
            return None, None
        return row.get(k), levels.get(f"items.{row_idx}.{k}")
    return None, None


def per_template_tables(docs: list[dict], lang: str = "zh") -> list[tuple]:
    groups: dict[int, list[dict]] = {}
    for d in docs:
        groups.setdefault(d["template_id"], []).append(d)
    tables = []
    for tid, gdocs in groups.items():
        tpl = db.get("templates", tid)
        if not tpl:
            continue
        sources = [(_col("batch", lang), "meta.batch"), (_col("file", lang), "meta.filename"),
                   (_col("pages", lang), "meta.pages")]
        sources += [(f.get("label") or f["key"], f"header.{f['key']}") for f in tpl["header_fields"]]
        sources += [(f.get("label") or f["key"], f"item.{f['key']}") for f in tpl["item_fields"]]
        columns = [c for c, _ in sources] + [_col("notes", lang)]
        rows = []
        for d in gdocs:
            rows += _doc_rows(d, tpl, [s for _, s in sources], lang)
        tables.append((tpl["name"], columns, rows))
    return tables


def _doc_rows(doc: dict, tpl: dict, sources: list[str], lang: str) -> list[list]:
    data = doc.get("data") or {"header": {}, "items": []}
    header, items = data.get("header", {}), data.get("items", [])
    levels = _issue_levels(doc)
    out = []
    for idx, row in (enumerate(items) if items else [(None, None)]):
        cells = [_resolve(s, doc, tpl, header, row, idx, levels) for s in sources]
        cells.append((_row_notes(doc, idx, lang), None))
        out.append(cells)
    return out


def unified_table(docs: list[dict], fmt: dict, lang: str = "zh") -> list[tuple]:
    columns = list(fmt["columns"])
    alloc = fmt.get("allocate") or {}
    do_alloc = alloc.get("enabled") and alloc.get("amount_col") in columns and alloc.get("method") in ("weight", "even")
    rows = []
    for d in docs:
        tpl = db.get("templates", d["template_id"])
        if not tpl:
            continue
        mapping = (fmt.get("mappings") or {}).get(str(tpl["id"]), {})
        doc_rows = _doc_rows(d, tpl, [mapping.get(c) or ("meta.batch" if c in BATCH_COLUMNS else "") for c in columns],
                             lang)
        if do_alloc:
            ai = columns.index(alloc["amount_col"])
            if (mapping.get(alloc["amount_col"]) or "").startswith("header.") and len(doc_rows) > 1:
                # A document-level amount repeated on every line would be counted
                # several times; keep it once so it gets spread over the lines.
                for r in doc_rows[1:]:
                    r[ai] = (None, r[ai][1])
            _allocate(doc_rows, columns, alloc, lang)
        rows += doc_rows
    out_cols = columns + ([_col("alloc", lang)] if do_alloc else []) + [_col("notes", lang)]
    if do_alloc:
        for r in rows:
            if len(r) == len(columns) + 1:  # rows without an allocation note
                r.insert(len(columns), ("", None))
    return [(fmt["name"], out_cols, rows)]


def _allocate(doc_rows: list[list], columns: list[str], alloc: dict, lang: str = "zh") -> None:
    """Spread a document's amount over rows that have none (e.g. one fee for several AWBs)."""
    ai = columns.index(alloc["amount_col"])
    amounts = [to_number(r[ai][0]) for r in doc_rows]
    if len(doc_rows) < 2 or all(a is not None for a in amounts) or not any(amounts):
        return
    total = round(sum(a for a in amounts if a), 2)
    method = alloc["method"]
    weights = None
    if method == "weight" and alloc.get("weight_col") in columns:
        wi = columns.index(alloc["weight_col"])
        weights = [to_number(r[wi][0]) for r in doc_rows]
        if any(w is None or w <= 0 for w in weights):
            weights = None
    if weights is None:
        weights, method = [1] * len(doc_rows), "even"
    wsum = sum(weights)
    shares = [round(total * w / wsum, 2) for w in weights]
    shares[-1] = round(total - sum(shares[:-1]), 2)
    if lang == "zh":
        note = f"{'按重量分摊' if method == 'weight' else '平均分摊'}，原金额合计 {total:,.2f}"
    else:
        note = f"{'Allocated by weight' if method == 'weight' else 'Allocated evenly'}, original total {total:,.2f}"
    for r, s in zip(doc_rows, shares):
        r[ai] = (s, r[ai][1])
        r.insert(len(columns), (note, None))


def _safe_sheet(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:28] or "Sheet"
    n, out = 2, base
    while out in used:
        out, n = f"{base}_{n}", n + 1
    used.add(out)
    return out


def to_xlsx(tables: list[tuple]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for title, columns, rows in tables:
        ws = wb.create_sheet(_safe_sheet(title, used))
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = FILL_HEAD
        for r in rows:
            ws.append([v for v, _ in r])
            for ci, (_, lvl) in enumerate(r, start=1):
                if lvl:
                    ws.cell(row=ws.max_row, column=ci).fill = FILL_ERR if lvl == "error" else FILL_WARN
        ws.freeze_panes = "A2"
        for ci, col in enumerate(columns, start=1):
            width = max([len(str(col))] + [len(str(r[ci - 1][0] or "")) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(ci)].width = min(max(width + 2, 8), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_bytes(columns, rows) -> bytes:
    s = io.StringIO()
    w = csv.writer(s)
    w.writerow(columns)
    for r in rows:
        w.writerow(["" if v is None else v for v, _ in r])
    return s.getvalue().encode("utf-8-sig")


def to_csv(tables: list[tuple]) -> tuple[bytes, str]:
    if len(tables) == 1:
        return _csv_bytes(tables[0][1], tables[0][2]), "csv"
    buf = io.BytesIO()
    used: set = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for title, columns, rows in tables:
            z.writestr(_safe_sheet(title, used) + ".csv", _csv_bytes(columns, rows))
    return buf.getvalue(), "zip"
