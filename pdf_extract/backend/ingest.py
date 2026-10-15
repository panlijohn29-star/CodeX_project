"""Turn uploads that are not PDFs into PDFs, so that template matching, extraction, review and export stay as they are.

Supported: PDF, images (jpg/png/webp/bmp/gif/tif), e-mails (.eml/.msg; attachments become documents of their own) and
spreadsheets (.xlsx/.xlsm/.csv, rendered as a table). Every conversion yields Items: the PDF to store, the name to show,
an optional note about where it came from, and a digest of the *original* content (used to recognise identical files).
"""
import csv
import hashlib
import html as htmllib
import io
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser

import pymupdf

from .i18n import bi

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
SHEET_EXTS = {".xlsx", ".xlsm", ".csv"}
MAIL_EXTS = {".eml", ".msg"}
SUPPORTED = {".pdf"} | IMAGE_EXTS | SHEET_EXTS | MAIL_EXTS
IGNORED_ATTACHMENTS = {".ics", ".vcf", ".p7s", ".asc", ".sig"}  # never invoices; not worth listing as "skipped"

MAX_FRAMES = 50          # pages taken from a multi-page TIFF / animated GIF
MAX_IMAGE_SIDE = 3000    # larger photos are scaled down; the AI only sees 1568 px anyway
MIN_IMAGE_BYTES = 15_000  # smaller images in an e-mail are signature logos, not documents
MAX_SHEET_ROWS = 300     # a longer table would not fit in one extraction result
MAX_PAGES = 40
MAX_MAIL_DEPTH = 2       # e-mails attached to e-mails

STORY_CSS = "body{font-family:sans-serif;font-size:10pt} td,th{padding:1pt 3pt;vertical-align:top} table{border-collapse:collapse}"


class IngestError(Exception):
    """msg is a bilingual {"zh", "en"} message."""

    def __init__(self, msg):
        self.msg = msg
        super().__init__(msg["en"])


@dataclass
class Item:
    filename: str       # name shown in the document list
    pdf: bytes
    source: dict | None  # bilingual note on where the document came from
    digest: str         # sha256 of the original content


def ext_of(name: str) -> str:
    return os.path.splitext(name or "")[1].lower()


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def convert(filename: str, data: bytes, depth: int = 0) -> list[Item]:
    ext = ext_of(filename)
    if ext == ".pdf":
        return [Item(filename, data, None, sha(data))]
    if ext in IMAGE_EXTS:
        return [Item(filename, _image_pdf(filename, data), None, sha(data))]
    if ext in SHEET_EXTS:
        return [_sheet_item(filename, data)]
    if ext in MAIL_EXTS:
        if depth > MAX_MAIL_DEPTH:
            raise IngestError(bi(f"邮件嵌套太深：{filename}", f"E-mails are nested too deeply: {filename}"))
        mail = _parse_eml(data) if ext == ".eml" else _parse_msg(filename, data)
        return _mail_items(filename, data, mail, depth)
    if ext == ".xls":
        raise IngestError(bi(f"不支持旧版 .xls：{filename}，请在 Excel 里另存为 .xlsx 再上传",
                             f"The old .xls format is not supported: {filename}. Save it as .xlsx in Excel and upload again"))
    raise IngestError(bi(f"不支持的文件类型：{filename}（支持 PDF、图片、邮件 .eml/.msg、表格 .xlsx/.csv）",
                         f"Unsupported file type: {filename} (supported: PDF, images, e-mail .eml/.msg, spreadsheets .xlsx/.csv)"))


# ------------------------------------------------------------------ rendering

def _story_pdf(html: str, landscape: bool = False, css: str = "") -> bytes:
    """Lay HTML out on A4 pages. The result has a real text layer, so it is handled like a digital PDF."""
    w, h = (842, 595) if landscape else (595, 842)
    story = pymupdf.Story(html=html, user_css=STORY_CSS + css)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    media = pymupdf.Rect(0, 0, w, h)
    more, pages = 1, 0
    while more and pages < MAX_PAGES:
        dev = writer.begin_page(media)
        more, _ = story.place(media + (30, 30, -30, -30))
        story.draw(dev)
        writer.end_page()
        pages += 1
    writer.close()
    return buf.getvalue()


def _esc(s) -> str:
    return htmllib.escape("" if s is None else str(s))


# --------------------------------------------------------------------- images

def _image_pdf(filename: str, data: bytes) -> bytes:
    from PIL import Image, ImageOps, ImageSequence

    try:
        im = Image.open(io.BytesIO(data))
        frames = []
        for n, frame in enumerate(ImageSequence.Iterator(im)):
            if n >= MAX_FRAMES:
                break
            fr = ImageOps.exif_transpose(frame)  # phone photos carry their rotation in EXIF
            if fr.mode in ("RGBA", "LA", "P"):
                rgba = fr.convert("RGBA")
                bg = Image.new("RGB", rgba.size, "white")
                bg.paste(rgba, mask=rgba.split()[3])
                fr = bg
            else:
                fr = fr.convert("RGB")
            fr.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            buf = io.BytesIO()
            fr.save(buf, "JPEG", quality=90)
            frames.append((fr.size, buf.getvalue()))
    except Exception as e:  # noqa: BLE001 - any decoder failure means the file is unusable
        raise IngestError(bi(f"无法打开图片：{filename}", f"Cannot open image: {filename}")) from e
    out = pymupdf.open()
    for (w, h), jpg in frames:
        page = out.new_page(width=w * 0.48, height=h * 0.48)  # 150 dpi, so rendering gives back the pixels 1:1
        page.insert_image(page.rect, stream=jpg)
    return out.tobytes()


# ----------------------------------------------------------------- spreadsheets

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.6f}".rstrip("0").rstrip(".")
    return " ".join(str(v).split())


def _clean(rows: list[list]) -> list[list[str]]:
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    width = max((max((i for i, c in enumerate(r) if c), default=-1) + 1 for r in rows), default=0)
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _read_csv(data: bytes) -> list[list]:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(filename: str, data: bytes) -> dict[str, list[list]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheets = {}
        for ws in wb.worksheets:
            if ws.sheet_state == "visible":
                sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return sheets
    except Exception as e:  # noqa: BLE001
        raise IngestError(bi(f"无法打开表格：{filename}", f"Cannot open spreadsheet: {filename}")) from e


def _sheet_item(filename: str, data: bytes) -> Item:
    raw = _read_csv(data) if ext_of(filename) == ".csv" else None
    sheets = {"CSV": raw} if raw is not None else _read_xlsx(filename, data)
    tables = {name: _clean(rows) for name, rows in sheets.items()}
    tables = {name: rows for name, rows in tables.items() if rows}
    if not tables:
        raise IngestError(bi(f"表格里没有内容：{filename}", f"The spreadsheet is empty: {filename}"))
    total = sum(len(r) for r in tables.values())
    if total > MAX_SHEET_ROWS:
        raise IngestError(bi(
            f"表格太长：{filename} 共 {total} 行，最多 {MAX_SHEET_ROWS} 行。请拆成几个文件上传",
            f"The spreadsheet is too long: {filename} has {total} rows, the limit is {MAX_SHEET_ROWS}. Split it into several files"))
    cols = max(len(r[0]) for r in tables.values())
    size = 9 if cols <= 7 else 8 if cols <= 10 else 7 if cols <= 14 else 6 if cols <= 20 else 5
    parts = []
    for name, rows in tables.items():
        body = "".join("<tr>" + "".join(f"<td>{_esc(c) or '&nbsp;'}</td>" for c in r) + "</tr>" for r in rows)
        parts.append(f"<p><b>{_esc(filename)}" + ("" if name == "CSV" else f" · {_esc(name)}") +
                     f"</b></p><table border='1'>{body}</table>")
    pdf = _story_pdf("".join(parts), landscape=cols > 7, css=f"td,th{{font-size:{size}pt}}")
    note = bi(f"由表格“{filename}”转换而来", f"Converted from the spreadsheet \"{filename}\"")
    return Item(filename, pdf, note, sha(data))


# ---------------------------------------------------------------------- e-mails

@dataclass
class Mail:
    subject: str
    sender: str
    to: str
    date: str
    html: str | None
    text: str | None
    attachments: list  # (filename, bytes, embedded_in_body)


def _parse_eml(data: bytes) -> Mail:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    body = msg.get_body(preferencelist=("html", "plain"))
    html = text = None
    if body is not None:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            html = content
        else:
            text = content
    attachments = []

    def walk(part):
        ctype = part.get_content_type()
        if ctype == "message/rfc822":
            inner = part.get_payload()[0]
            attachments.append((part.get_filename() or "forwarded.eml", inner.as_bytes(), False))
        elif part.is_multipart():
            for p in part.iter_parts():
                walk(p)
        elif part is not body:
            name, disp = part.get_filename(), part.get_content_disposition()
            if ctype.startswith("text/") and not name and disp != "attachment":
                return  # the other alternative of the body
            payload = part.get_payload(decode=True) or b""
            if not name and ctype.startswith("image/"):
                name = "image." + ctype.split("/")[1].replace("jpeg", "jpg")
            if name:
                attachments.append((name, payload, disp == "inline" or (part.get("Content-ID") and not disp)))

    walk(msg)
    return Mail(str(msg.get("subject") or ""), str(msg.get("from") or ""), str(msg.get("to") or ""),
                str(msg.get("date") or ""), html, text, attachments)


def _parse_msg(filename: str, data: bytes) -> Mail:
    try:
        import extract_msg
    except ImportError as e:
        raise IngestError(bi("读取 .msg 需要安装 extract-msg（pip install extract-msg）",
                             ".msg files need the extract-msg package (pip install extract-msg)")) from e
    fd, tmp = tempfile.mkstemp(suffix=".msg")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        m = extract_msg.openMsg(tmp)
        try:
            html = m.htmlBody.decode("utf-8", "replace") if isinstance(m.htmlBody, (bytes, bytearray)) else m.htmlBody
            attachments = []
            for a in m.attachments:
                payload = a.data
                name = a.longFilename or a.shortFilename or ""
                if isinstance(payload, (bytes, bytearray)) and name:
                    attachments.append((name, bytes(payload), bool(getattr(a, "cid", None))))
            return Mail(str(m.subject or ""), str(m.sender or ""), str(m.to or ""), str(m.date or ""),
                        html or None, m.body or None, attachments)
        finally:
            m.close()
    except IngestError:
        raise
    except Exception as e:  # noqa: BLE001
        raise IngestError(bi(f"无法打开邮件：{filename}", f"Cannot open e-mail: {filename}")) from e


def _clean_html(html: str) -> str:
    html = re.sub(r"<!--.*?-->", "", html, flags=re.S)
    html = re.sub(r"<(script|style|head|title)\b.*?</\1\s*>", "", html, flags=re.S | re.I)
    html = re.sub(r"<(img|link|meta)\b[^>]*>", "", html, flags=re.I)
    m = re.search(r"<body\b[^>]*>(.*)</body\s*>", html, flags=re.S | re.I)
    return m.group(1) if m else html


def _body_pdf(mail: Mail) -> bytes | None:
    """The mail's headers and body as a PDF; None when there is no body text."""
    if mail.html and re.sub(r"<[^>]+>|&nbsp;|\s", "", _clean_html(mail.html)):
        body = _clean_html(mail.html)
    elif mail.text and mail.text.strip():
        body = _esc(mail.text.strip()).replace("\n", "<br>")
    else:
        return None
    head = (f"<p><b>Subject:</b> {_esc(mail.subject)}<br><b>From:</b> {_esc(mail.sender)}<br>"
            f"<b>To:</b> {_esc(mail.to)}<br><b>Date:</b> {_esc(mail.date)}</p><hr>")
    return _story_pdf(head + body)


def _join(*msgs) -> dict | None:
    msgs = [m for m in msgs if m]
    return bi(" · ".join(m["zh"] for m in msgs), " · ".join(m["en"] for m in msgs)) if msgs else None


def _mail_items(filename: str, raw: bytes, mail: Mail, depth: int) -> list[Item]:
    usable: list[Item] = []
    skipped: list[str] = []
    for name, payload, embedded in mail.attachments:
        ext = ext_of(name)
        if ext in IGNORED_ATTACHMENTS:
            continue
        if ext not in SUPPORTED:
            skipped.append(name)
            continue
        if ext in IMAGE_EXTS and (embedded or len(payload) < MIN_IMAGE_BYTES):
            continue  # a picture inside the body or a logo
        try:
            usable += convert(name, payload, depth + 1)
        except IngestError:
            skipped.append(name)
    who = " · ".join(x for x in (mail.sender, mail.date) if x)
    note = bi(f"邮件“{mail.subject or filename}”" + (f"（{who}）" if who else ""),
              f"E-mail \"{mail.subject or filename}\"" + (f" ({who})" if who else ""))
    miss = bi(f"未处理的附件：{', '.join(skipped)}", f"Attachments not processed: {', '.join(skipped)}") if skipped else None
    if usable:
        # An attachment's own note (e.g. converted from a spreadsheet) is kept after the mail's.
        return [Item(i.filename, i.pdf, _join(note, i.source, miss), i.digest) for i in usable]
    pdf = _body_pdf(mail)
    if pdf is None:
        raise IngestError(bi(f"邮件里没有可处理的附件或正文：{filename}",
                             f"The e-mail has no attachment or body that can be processed: {filename}"))
    body_note = bi(f"邮件正文：{mail.subject or filename}" + (f"（{who}）" if who else ""),
                   f"Mail body: {mail.subject or filename}" + (f" ({who})" if who else ""))
    return [Item(filename, pdf, _join(body_note, miss), sha(raw))]
