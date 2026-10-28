from __future__ import annotations

import os
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from features.invoice_lookup_core import (ALLOWED_CREATORS, SPECIAL_CUSTOMER, IWT_CUSTOMERS, LookupError, RawRow, Snapshot,
                                          Source, decimal_value, evaluate, fetch_snapshot, key, load_db_config,
                                          normalized_do, parse_since, query_keys, read_source, source_sheets)
from features.invoice_lookup_excel import export_result, make_workbook
from openpyxl import Workbook, load_workbook

_TMP = tempfile.TemporaryDirectory()
SOURCE_FILE = Path(_TMP.name) / "raw.xlsx"
_wb = Workbook()
_wb.active.title = "raw"
_wb.active.append(["PRO", "AMT", "DO#"])
_wb.save(SOURCE_FILE)


def raw(pro="P1", amount="100", do="DO1", customer=None, ap=None, line=2):
    values = [pro, amount, do, ap, None, customer, None]
    return RawRow(line, values, str(pro), normalized_do(do), decimal_value(amount))


def invoice(pro="P1", amount="100", number="INV1", customer="CUSTOMER A", creator="USSDC.ORD"):
    return dict(customer_invoice_no=pro, amount_balance=amount, invoice_no=number,
                customer=customer, invoice_by=creator, invoice_date="2026-09-21 00:00:00", balance="VENDOR")


def case(rows=None, invoices=None, customers=None, xc=None, sc=None):
    rows = rows if rows is not None else [raw()]
    source = Source(SOURCE_FILE, "raw", rows)
    dos, pros = query_keys(source)
    if xc is None:
        customer_records = customers if customers is not None else [dict(dispatch_no=d, customer_code="CUSTOMER A") for d in dos]
        xc = Snapshot(customer_records, set(dos))
    if sc is None:
        sc = Snapshot(invoices if invoices is not None else [invoice()], set(pros))
    return evaluate(source, xc, sc)


def codes(result):
    return {i.code for i in result.issues}


class MatchingTests(unittest.TestCase):
    def test_repeated_pro_and_repeated_invoice_not_double_counted(self):
        result = case([raw(line=2), raw(do="DO2", line=3)])
        self.assertEqual([r.selected for r in result.source.rows], ["INV1", "INV1"])
        self.assertEqual(result.reconciliations[0].actual, Decimal("100.00"))
        self.assertEqual(result.reconciliations[0].status, "PASS")

    def test_special_customer_split(self):
        result = case([raw(line=2), raw(do="DO2", line=3)],
                      [invoice(amount="70"), invoice(amount="30", number="INV2", customer=SPECIAL_CUSTOMER, creator="SHILONG.ORD")],
                      [dict(dispatch_no="DO1", customer_code="CUSTOMER B"), dict(dispatch_no="DO2", customer_code=SPECIAL_CUSTOMER)])
        self.assertEqual([r.selected for r in result.source.rows], ["INV1", "INV2"])
        self.assertEqual(result.reconciliations[0].status, "PASS")
        self.assertEqual(result.status, "PASS")
        self.assertEqual(result.issues, [])
        wb = make_workbook(result)
        try:
            for name in ("Raw_INV", "Look_UP"):
                for cells in wb[name].iter_rows(min_row=2):
                    self.assertTrue(all(c.fill.patternType is None for c in cells))
        finally:
            wb.close()

    def test_special_customer_split_mismatch_still_highlights_both_sheets(self):
        result = case([raw(line=2), raw(do="DO2", line=3)],
                      [invoice(amount="70"), invoice(amount="29.99", number="INV2", customer=SPECIAL_CUSTOMER)],
                      [dict(dispatch_no="DO1", customer_code="A"), dict(dispatch_no="DO2", customer_code=SPECIAL_CUSTOMER)])
        self.assertEqual([r.selected for r in result.source.rows], ["INV1", "INV2"])
        self.assertEqual(codes(result), {"AMOUNT_MISMATCH"})
        self.assertEqual(result.reconciliations[0].difference, Decimal("-0.01"))
        self.assertEqual(result.reconciliations[0].status, "FAIL")
        wb = make_workbook(result)
        try:
            for name in ("Raw_INV", "Look_UP"):
                for cells in wb[name].iter_rows(min_row=2):
                    self.assertTrue(all(c.fill.fgColor.rgb == "00FCE4E4" for c in cells))
        finally:
            wb.close()

    def test_non_special_never_uses_special_ticket(self):
        result = case(invoices=[invoice(customer=SPECIAL_CUSTOMER)])
        self.assertEqual(result.source.rows[0].selected, "")
        self.assertIn("NO_CUSTOMER_INVOICE", codes(result))

    def test_nonallowed_only_is_blank(self):
        result = case(invoices=[invoice(creator="OTHER")])
        self.assertEqual(result.source.rows[0].selected, "")
        self.assertIn("NO_ALLOWED_INVOICE", codes(result))

    def test_equal_amount_outside_creator_warns_but_does_not_count(self):
        result = case(invoices=[invoice(), invoice(number="INV2", creator="OTHER")])
        self.assertEqual(result.source.rows[0].selected, "INV1")
        self.assertEqual(result.reconciliations[0].actual, Decimal("100"))
        self.assertIn("SUSPECT_DUPLICATE_INVOICE", codes(result))

    def test_multiple_qualified_blank_even_when_sum_matches(self):
        result = case(invoices=[invoice(amount=40), invoice(number="INV2", amount=60, creator="JOHNWU.ORD")])
        self.assertEqual(result.source.rows[0].selected, "")
        self.assertIn("AMBIGUOUS_INVOICE", codes(result))
        self.assertEqual(result.reconciliations[0].status, "INCOMPLETE")

    def test_iwt_unique_designated_customer_and_ordinary_invoices_excluded(self):
        single = case([raw(do="IWT1")], [invoice(customer="USA-IWT")])
        self.assertEqual(single.source.rows[0].selected, "INV1")
        self.assertEqual(single.source.rows[0].customer, "")
        self.assertEqual(single.reconciliations[0].status, "PASS")
        self.assertEqual(single.status, "PASS")
        self.assertEqual(single.issues, [])
        split = case([raw(do="IWT1")], [invoice(amount=70), invoice(number="INV2", amount=30, customer=SPECIAL_CUSTOMER)])
        self.assertEqual(split.source.rows[0].selected, "")
        self.assertIn("NO_CUSTOMER_INVOICE", codes(split))

    def test_incomplete_iwt_prevents_pass_even_if_other_rows_total_matches(self):
        rows = [raw(line=2), raw(do="DO2", line=3), raw(do="IWT1", line=4)]
        result = case(rows, [invoice(amount=70), invoice(number="INV2", amount=30, customer=SPECIAL_CUSTOMER)],
                      [dict(dispatch_no="DO1", customer_code="A"), dict(dispatch_no="DO2", customer_code=SPECIAL_CUSTOMER)])
        self.assertEqual(result.reconciliations[0].actual, Decimal("100"))
        self.assertEqual(result.reconciliations[0].status, "INCOMPLETE")
        self.assertIsNone(result.reconciliations[0].difference)

    def test_raw_amount_conflict_is_undetermined(self):
        result = case([raw(line=2), raw(amount="90", do="DO2", line=3)])
        self.assertIn("RAW_AMOUNT_CONFLICT", codes(result))
        self.assertIsNone(result.reconciliations[0].expected)
        self.assertEqual(result.reconciliations[0].status, "UNDETERMINED")

    def test_one_cent_is_fail(self):
        result = case(invoices=[invoice(amount="100.01")])
        self.assertEqual(result.reconciliations[0].difference, Decimal("0.01"))
        self.assertIn("AMOUNT_MISMATCH", codes(result))

    def test_rounding_and_negative_amounts(self):
        result = case([raw(amount="-100.005")], [invoice(amount="-100.005")])
        self.assertEqual(result.reconciliations[0].expected, Decimal("-100.01"))
        self.assertEqual(result.reconciliations[0].status, "PASS")
        self.assertTrue({"AMOUNT_PRECISION", "INVOICE_PRECISION", "NONPOSITIVE_AMOUNT"}.issubset(codes(result)))

    def test_zero_is_not_missing(self):
        result = case([raw(amount=0)], [invoice(amount=0)])
        self.assertEqual(result.reconciliations[0].actual, Decimal("0.00"))
        self.assertEqual(result.reconciliations[0].status, "PASS")

    def test_ar_excluded(self):
        result = case([raw(pro="DO10 - AR")], [])
        self.assertEqual(query_keys(result.source)[1], [])
        self.assertEqual(result.reconciliations[0].status, "EXCLUDED_AR")
        self.assertNotIn("NO_SC_INVOICE", codes(result))

    def test_customer_missing_conflict_and_stale_values(self):
        missing = case([raw(customer="OLD", ap="OLDINV")], customers=[])
        self.assertIn("CUSTOMER_MISSING", codes(missing))
        self.assertIn("CUSTOMER_CHANGED", codes(missing))
        self.assertIn("AP_INVOICE_CHANGED", codes(missing))
        self.assertEqual(missing.source.rows[0].customer, "")
        self.assertEqual(missing.reconciliations[0].status, "INCOMPLETE")
        conflict = case(customers=[dict(dispatch_no="DO1", customer_code="A"), dict(dispatch_no="DO1", customer_code="B")])
        self.assertIn("CUSTOMER_CONFLICT", codes(conflict))

    def test_no_invoice_and_failed_query_are_different(self):
        absent = case(invoices=[])
        failed = case(sc=Snapshot([], set(), {"P1": "SC 查询失败"}))
        self.assertIn("NO_SC_INVOICE", codes(absent))
        self.assertIn("SC_QUERY_INCOMPLETE", codes(failed))
        self.assertNotIn("NO_SC_INVOICE", codes(failed))

    def test_identical_db_records_count_once_and_conflicts_block(self):
        repeated = case(invoices=[invoice(), invoice()])
        self.assertIn("DUPLICATE_DB_ROW", codes(repeated))
        self.assertEqual(repeated.reconciliations[0].actual, Decimal("100"))
        conflicting = case(invoices=[invoice(), invoice(amount=90)])
        self.assertIn("INVOICE_CONFLICT", codes(conflicting))
        self.assertFalse(conflicting.source.rows[0].selected)

    def test_invoice_cross_pro_blocks_both(self):
        result = case([raw(line=2), raw(pro="P2", do="DO2", line=3)], [invoice(), invoice(pro="P2")])
        self.assertIn("INVOICE_CROSS_PRO", codes(result))
        self.assertTrue(all(not r.selected for r in result.source.rows))

    def test_invalid_candidate_must_not_disappear(self):
        result = case(invoices=[invoice(customer=None), invoice(number="INV2")])
        self.assertIn("BLOCKED_CANDIDATE", codes(result))
        self.assertFalse(result.source.rows[0].selected)

    def test_duplicate_source_and_normalized_do_are_retained(self):
        result = case([raw(line=2), raw(line=3), raw(do="DO1 - 1", line=4)])
        self.assertEqual(len(result.source.rows), 3)
        self.assertTrue({"DUPLICATE_RAW_ROW", "REPEATED_DO"}.issubset(codes(result)))


class InputAndExportTests(unittest.TestCase):
    def test_do_suffix_only_exact_one(self):
        for value in ["DO123-1", "DO123- 1", "DO123 -1", "DO123 - 1"]:
            self.assertEqual(normalized_do(value), "DO123")
        for value in ["DO123-11", "DO123-2", "IWT123-1", "DO123-10"]:
            self.assertEqual(normalized_do(value), value)

    def test_dates_and_invalid_amounts(self):
        for value in [None, True, "bad", "NaN", "Infinity"]:
            self.assertIsNone(decimal_value(value))
        for value in ["2026-02-30", "9/1/2026", "2026-9-1"]:
            with self.assertRaises(LookupError):
                parse_since(value)

    def test_source_selection_leading_zeros_invalid_rows_and_formula_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "one"
            ws.append(["PRO#", "AMT", "DO#", "remark"])
            ws.append([123, 10, "DO123 -1", "=literal"])
            ws["A2"].number_format = "00000"
            ws.append(["00100-02", "=1+2", "DO2", ""])
            ws.append([None, "bad", "DO3", "invalid"])
            ws2 = wb.create_sheet("two")
            ws2.append(["PRO", "AMT", "DO#"])
            ws2.append(["P", 1, "DO4"])
            wb.save(path)
            self.assertEqual(source_sheets(path), ["one", "two"])
            with self.assertRaises(LookupError):
                read_source(path)
            source = read_source(path, "one")
            self.assertEqual(len(source.rows), 3)
            self.assertEqual(source.rows[0].pro, "00123")
            self.assertEqual(source.rows[1].pro, "00100-02")
            self.assertTrue(source.rows[1].input_errors)
            self.assertTrue(source.rows[2].input_errors)

    def test_export_preserves_source_and_no_formula_injection(self):
        result = case([raw(line=2), raw(do="DO2", line=3)])
        result.source.rows[0].values[6] = '=HYPERLINK("https://example.invalid","x")'
        with tempfile.TemporaryDirectory() as tmp:
            first = export_result(result, Path(tmp))
            second = export_result(result, Path(tmp))
            self.assertNotEqual(first, second)
            wb = load_workbook(first)
            self.assertEqual(wb.sheetnames, ["Raw_INV", "Look_UP", "Validation_Report"])
            self.assertEqual(wb["Raw_INV"].max_row, 3)
            self.assertEqual(wb["Raw_INV"]["G2"].data_type, "s")
            self.assertEqual(wb["Raw_INV"]["D2"].value, "INV1")
            self.assertEqual(wb["Raw_INV"].freeze_panes, "D2")
            self.assertEqual(wb["Raw_INV"].tables["RawInvoices"].ref, "A1:K3")
            self.assertTrue(wb["Validation_Report"].tables)
            self.assertTrue(any(c.hyperlink for row in wb["Validation_Report"] for c in row))
            self.assertFalse(any(c.data_type == "f" for ws in wb for row in ws for c in row))
            wb.close()


class MixedDoIwtTests(unittest.TestCase):
    def mixed_25427(self, iwt_customer="USA-IWT", iwt_amount="2296.49"):
        return case(
            [raw("25427", "4100", "DO202609131031", line=46),
             raw("25427", "4100", "DO202609131032", line=47),
             raw("25427", "4100", "IWT202609131002", line=48)],
            [invoice("25427", "1803.51", "ORDAP260914827", "TESLA GROUP"),
             invoice("25427", iwt_amount, "ORDAP260916077", iwt_customer)],
            [dict(dispatch_no="DO202609131031", customer_code="TESLA GROUP"),
             dict(dispatch_no="DO202609131032", customer_code="TESLA GROUP")],
        )

    def test_pro_25427_both_iwt_customer_names(self):
        for name in IWT_CUSTOMERS:
            with self.subTest(customer=name):
                result = self.mixed_25427(name)
                self.assertEqual([r.selected for r in result.source.rows],
                                 ["ORDAP260914827", "ORDAP260914827", "ORDAP260916077"])
                recon = result.reconciliations[0]
                self.assertEqual(recon.expected, Decimal("4100.00"))
                self.assertEqual(recon.actual, Decimal("4100.00"))
                self.assertEqual(recon.status, "PASS")
                self.assertEqual(result.status, "PASS")
                self.assertEqual(result.issues, [])
                wb = make_workbook(result)
                try:
                    self.assertEqual([wb["Raw_INV"].cell(r, 4).value for r in (2, 3, 4)],
                                     ["ORDAP260914827", "ORDAP260914827", "ORDAP260916077"])
                    self.assertIn("IWT", wb["Look_UP"]["I3"].value)
                    for name, status_column in (("Raw_INV", 9), ("Look_UP", 10)):
                        for cells in wb[name].iter_rows(min_row=2):
                            self.assertEqual(cells[status_column - 1].value, "PASS")
                            self.assertTrue(all(c.fill.patternType is None for c in cells))
                finally:
                    wb.close()

    def test_iwt_customer_names_trim_and_case(self):
        result = self.mixed_25427("  apex logistics international (ord), inc.  ")
        self.assertEqual(result.source.rows[-1].selected, "ORDAP260916077")

    def test_mixed_amount_mismatch_is_still_reported(self):
        result = self.mixed_25427(iwt_amount="2296.48")
        self.assertIn("AMOUNT_MISMATCH", codes(result))
        self.assertEqual(result.reconciliations[0].difference, Decimal("-0.01"))

    def test_two_iwt_invoices_remain_ambiguous(self):
        result = case([raw(do="IWT1")],
                      [invoice(amount="70", customer="USA-IWT"),
                       invoice(amount="30", number="INV2", customer="APEX LOGISTICS INTERNATIONAL (ORD), INC.")])
        self.assertEqual(result.source.rows[0].selected, "")
        self.assertIn("AMBIGUOUS_INVOICE", codes(result))
        self.assertEqual(result.reconciliations[0].status, "INCOMPLETE")

    def test_iwt_creator_filter_never_falls_back_to_do(self):
        result = case([raw(do="IWT1")],
                      [invoice(), invoice(number="INV2", customer="USA-IWT", creator="OTHER")])
        self.assertEqual(result.source.rows[0].selected, "")
        self.assertIn("NO_ALLOWED_INVOICE", codes(result))

    def test_do_excludes_unique_iwt_invoice_even_if_customer_is_unknown(self):
        for customers in (None, []):
            with self.subTest(customers=customers):
                result = case(invoices=[invoice(customer="USA-IWT")], customers=customers)
                self.assertEqual(result.source.rows[0].selected, "")
                self.assertIn("NO_CUSTOMER_INVOICE", codes(result))

    def test_do_exact_customer_disambiguates_ordinary_invoices(self):
        result = case([raw(line=2), raw(do="DO2", line=3)],
                      [invoice(amount="70", customer="A"), invoice(amount="30", number="INV2", customer="B")],
                      [dict(dispatch_no="DO1", customer_code="A"), dict(dispatch_no="DO2", customer_code="B")])
        self.assertEqual([r.selected for r in result.source.rows], ["INV1", "INV2"])
        self.assertEqual(result.reconciliations[0].status, "PASS")

    def test_do_shared_invoice_fallback_preserves_creator_preference(self):
        result = case(invoices=[invoice(creator="OTHER"), invoice(number="INV2", customer="B")])
        self.assertEqual(result.source.rows[0].selected, "INV2")
        self.assertEqual(result.reconciliations[0].status, "PASS")

    def test_three_way_ordinary_kn_and_iwt_split(self):
        result = case([raw(line=2), raw(do="DO2", line=3), raw(do="IWT1", line=4)],
                      [invoice(amount="50"), invoice(amount="20", number="KN", customer=SPECIAL_CUSTOMER),
                       invoice(amount="30", number="IWT", customer="USA-IWT")],
                      [dict(dispatch_no="DO1", customer_code="CUSTOMER A"), dict(dispatch_no="DO2", customer_code=SPECIAL_CUSTOMER)])
        self.assertEqual([r.selected for r in result.source.rows], ["INV1", "KN", "IWT"])
        self.assertEqual(result.reconciliations[0].status, "PASS")
        self.assertNotIn("MULTIPLE_PRO_INVOICES", codes(result))



class DatabaseTests(unittest.TestCase):
    def test_sc_uses_platform_scdbus_profile(self):
        env = {"DB_HOST": "h", "DB_USER": "u", "DB_PASSWORD": " pw<x> ", "SCDBUS_DATABASE": "sc"}
        with patch.dict("os.environ", env, clear=True):
            config = load_db_config(None, "SC")
        self.assertEqual((config["host"], config["db"], config["password"]), ("h", "sc", " pw<x> "))
        self.assertNotIn("ssl", config)

    def test_xc_uses_dedicated_settings_and_optional_tls(self):
        env = {"XC_DB_HOST": "x", "XC_DB_USER": "u", "XC_DB_PASSWORD": "p", "XC_DB_NAME": "xc", "XC_DB_SSL": "true",
               "SYSTEMROOT": os.environ.get("SYSTEMROOT", "")}
        with patch.dict("os.environ", env, clear=True):
            config = load_db_config(None, "XC")
        self.assertEqual(config["db"], "xc")
        self.assertTrue(config["ssl"].check_hostname)

    def test_missing_config_is_a_safe_lookup_error(self):
        with patch.dict("os.environ", {}, clear=True):
            with self.assertRaises(LookupError):
                load_db_config(None, "XC")
            with self.assertRaises(LookupError):
                load_db_config(None, "SC")

    def test_parameterized_batches_and_failed_batch_preserve_success(self):
        calls = []
        class Cursor:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def execute(self, sql, params):
                calls.append((sql, params))
                if len(calls) == 2:
                    raise Exception(2013, "secret-password-must-not-appear")
            def fetchall(self): return [{"CUSTOMER_INVOICE_NO": "P0", "INVOICE_NO": "INV"}]
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        keys = [f"P{i}" for i in range(501)]
        keys[4] = "P' OR 1=1 --"
        snapshot = fetch_snapshot("SC", keys, "2026-09-01", None, connector=lambda **_: Connection(), config_loader=lambda *_: {})
        self.assertEqual(len(snapshot.succeeded), 500)
        self.assertEqual(len(snapshot.failed), 1)
        self.assertNotIn(keys[4], calls[0][0])
        self.assertIn(keys[4], calls[0][1])
        self.assertEqual(calls[0][1][-2:], ("2026-09-01", "3"))
        self.assertNotIn("secret", str(snapshot.failed))


if __name__ == "__main__":
    unittest.main(verbosity=2)

