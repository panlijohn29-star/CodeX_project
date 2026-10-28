import ast
import contextlib
import importlib.util
import io
from decimal import Decimal as D
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import features.ord_do_allocation as m
import features.ord_do_allocation_language as language
from openpyxl import load_workbook


def raw(do, amount=None, remark='', total='4800', pro='P'):
    assigned, invalid = m.parse_assigned_amount(remark, do)
    return m.RawRow(2, do, pro, D(total), remark=remark,
                    assigned_amount=D(amount) if amount is not None else assigned,
                    remark_invalid=invalid)


def calc(do, weight, hbl='H', customer='TEST'):
    return m.CalcRow(do, hbl, D(weight), D(weight), 'DO', sc_job_no='JOB', customer=customer)


def run(sources, shipments):
    v = m.ValidationLog()
    m.apply_raw_mappings(shipments, sources)
    m.allocate_costs(shipments, v, sources)
    return v


class AllocationTests(unittest.TestCase):
    def test_formats(self):
        for text in ['PLEASE ALLOCATE 3300USD TO THIS DO', 'pls allocate 3300to this do (DO1)',
                     'allocate $3,300.00 to DO1', 'please allocate USD 3300 to this do.']:
            self.assertEqual(m.parse_assigned_amount(text, 'DO1'), (D('3300'), False))
        self.assertEqual(m.parse_assigned_amount('PLEASE ALLOCATE 3,013.06USD TO THIS DO', 'DO1'), (D('3013.06'), False))

    def test_ambiguous_remarks_are_not_applied(self):
        for text in ['do not allocate 3300 to this do', 'allocate 3300 or 1500 to this do',
                     'allocate -10 to this do', 'allocate 33,00 to this do',
                     'allocate 3.333 to this do', 'allocate 3300 to DO2']:
            self.assertEqual(m.parse_assigned_amount(text, 'DO1'), (None, True))
        self.assertEqual(m.parse_assigned_amount('delivered yesterday', 'DO1'), (None, False))

    def test_full_override_splits_within_do(self):
        shipments = [calc('DO1','1','H1'),calc('DO1','2','H2'),calc('DO2','50','H3')]
        v = run([raw('DO1','3300'),raw('DO2','1500')], shipments)
        self.assertEqual([r.actual_cost for r in shipments], [D('1100'), D('2200'), D('1500')])
        self.assertEqual(v.counts(), {'INFO':2})

    def test_partial_override(self):
        shipments=[calc('DO1','100'),calc('DO2','1'),calc('DO3','2')]
        run([raw('DO1','3300'),raw('DO2'),raw('DO3')],shipments)
        self.assertEqual([r.actual_cost for r in shipments],[D('3300'),D('500'),D('1000')])

    def test_no_remark_legacy(self):
        shipments=[calc('DO1','1'),calc('DO2','9')]
        run([raw('DO1'),raw('DO2')],shipments)
        self.assertEqual([r.actual_cost for r in shipments],[D('480'),D('4320')])
        self.assertTrue(all(not r.allocation_group for r in shipments))

    def test_ordinary_remark_legacy(self):
        shipments=[calc('DO1','1'),calc('DO2','9')]
        run([raw('DO1',remark='delivered'),raw('DO2')],shipments)
        self.assertEqual(shipments[0].actual_cost,D('480'))

    def test_over_under_and_conflicting_assignments(self):
        for sources in [[raw('DO1','4000'),raw('DO2','1000')],
                        [raw('DO1','3000'),raw('DO2','1000')],
                        [raw('DO1','3300'),raw('DO1','3200'),raw('DO2')],
                        [raw('DO1',remark='allocate ??'),raw('DO2')]]:
            shipments=[calc('DO1','1'),calc('DO2','1')]
            v=run(sources,shipments)
            self.assertIn('REMARK_ALLOCATION_CONFLICT',[i.issue_type for i in v.items])
            self.assertTrue(all(r.allocation_invalid and r.actual_cost is None for r in shipments))

    def test_zero_assignment(self):
        shipments=[calc('DO1','1'),calc('DO2','1')]
        run([raw('DO1','0'),raw('DO2')],shipments)
        self.assertEqual([r.actual_cost for r in shipments],[D('0'),D('4800')])

    def test_missing_target_does_not_redistribute_it(self):
        shipments=[calc('DO2','1')]
        v=run([raw('DO1','3300'),raw('DO2')],shipments)
        self.assertEqual(shipments[0].actual_cost,D('1500'))
        self.assertIn('REMARK_ALLOCATION_UNFULFILLED',[i.issue_type for i in v.items])

    def test_rounding_and_invalid_cw(self):
        shipments=[calc('DO1','0'),calc('DO1','1'),calc('DO1','1'),calc('DO1','1')]
        v=run([raw('DO1','100',total='100')],shipments)
        self.assertEqual([r.actual_cost for r in shipments],[D('0'),D('33.34'),D('33.33'),D('33.33')])
        self.assertIn('INVALID_H_CW',[i.issue_type for i in v.items])

    def test_shared_customer_stays_pro_level(self):
        shipments=[calc('DO1','1',customer='VERTIV CORPORATION'),calc('DO2','1',customer='OTHER')]
        run([raw('DO1','3300'),raw('DO2','1500')],shipments)
        self.assertTrue(all(r.share_other_customer=='Y' for r in shipments))

    def test_export_formulas_and_repeated_job(self):
        sources=[raw('DO1','3300'),raw('DO2','1500')]
        shipments=[calc('DO1','1',customer='CORNING INCORPORATED'),calc('DO2','9',customer='CORNING INCORPORATED')]
        v=run(sources,shipments)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'test.xlsx'
            m.write_workbook(path,sources,shipments,{},m.CarrierLookup({},{}),{},v)
            w=load_workbook(path)
            self.assertIn('P2',w['calculation']['H2'].value)
            self.assertIn('Q2',w['calculation']['I2'].value)
            self.assertEqual(w['DO_ARupload']['K2'].value,"='calculation'!I2")
            self.assertEqual(w['DO_ARupload']['K3'].value,"='calculation'!I3")
            self.assertEqual(w['DO_upload']['K2'].value,3300)
            self.assertEqual(w['DO_upload']['K3'].value,1500)
            self.assertEqual([c.value for c in w['DO_upload'][1]],m.UPLOAD_HEADERS)
            w.close()

    def test_bilingual_validation_and_progress(self):
        for lang, needle in [('zh','备注指定金额'),('en','Remark-based allocation')]:
            language.set_language(lang,persist=False)
            v=run([raw('DO1','4800')],[calc('DO1','1')])
            self.assertIn(needle,v.items[0].message)
            stream=io.StringIO()
            with contextlib.redirect_stdout(stream): m.progress('Created: test.xlsx')
            self.assertIn('已生成' if lang=='zh' else 'Created',stream.getvalue())
        tree=ast.parse(Path(m.__file__).read_text(encoding='utf-8'))
        codes={node.args[2].value for node in ast.walk(tree) if isinstance(node,ast.Call)
               and isinstance(node.func,ast.Attribute) and node.func.attr=='add'
               and len(node.args)>3 and isinstance(node.args[2],ast.Constant)
               and isinstance(node.args[0],ast.Constant) and node.args[0].value in ('ERROR','WARNING','INFO')}
        self.assertFalse(codes-language.VALIDATION_ZH.keys())

    def test_bilingual_exports_keep_business_values(self):
        sources=[raw('DO1','3300'),raw('DO2','1500')]
        shipments=[calc('DO1','1'),calc('DO2','1')]
        snapshots=[]
        with tempfile.TemporaryDirectory() as directory:
            for lang in ('zh','en'):
                language.set_language(lang,persist=False)
                v=run(sources,shipments)
                path=Path(directory)/f'{lang}.xlsx'
                m.write_workbook(path,sources,shipments,{},m.CarrierLookup({},{}),{},v)
                w=load_workbook(path)
                snapshots.append({s.title:list(s.values) for s in w})
                w.close()
        for sheet in snapshots[0]:
            if sheet!='Validation_report': self.assertEqual(snapshots[0][sheet],snapshots[1][sheet])
        self.assertNotEqual(snapshots[0]['Validation_report'][1][7],snapshots[1]['Validation_report'][1][7])


if __name__=='__main__':
    unittest.main()
