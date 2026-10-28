import ast
import inspect
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import features.ord_do_allocation as m
import features.ord_do_allocation_language as language


class WmsTests(unittest.TestCase):
    def test_config_from_platform_environment(self):
        env = {'XC_DB_HOST': 'x', 'XC_DB_USER': 'u', 'XC_DB_PASSWORD': 'p', 'XC_DB_NAME': 'xc',
               'WMS_DB_HOST': 'w', 'WMS_DB_USER': 'u', 'WMS_DB_PASSWORD': 'p', 'WMS_DB_NAME': 'wms',
               'DB_HOST': 'sc', 'DB_USER': 'u', 'DB_PASSWORD': 'p', 'SCDBUS_DATABASE': 'scdbus'}
        with patch.dict(os.environ, env, clear=True):
            xc, sc, wms = m.load_config()
        self.assertEqual((xc['database'], sc['db'], wms['database']), ('xc', 'scdbus', 'wms'))
        self.assertNotIn('ssl', xc)

    def test_missing_wms_configuration_is_reported(self):
        env = {'XC_DB_HOST': 'x', 'XC_DB_USER': 'u', 'XC_DB_PASSWORD': 'p', 'XC_DB_NAME': 'xc'}
        with patch.dict(os.environ, env, clear=True):
            with self.assertRaises(m.ConverterError) as caught:
                m.load_config()
        self.assertIn('WMS_DB_HOST', str(caught.exception))

    def test_all_query_routes(self):
        tree = ast.parse(inspect.getsource(m.build_workbook))
        expected = {'query_iwt_meta': 'wms_config', 'query_iwt_calc': 'wms_config',
                    'query_iwt_to': 'wms_config', 'query_dispatch_meta': 'xc_config',
                    'query_xc_dispatch_calc': 'xc_config', 'query_sc_fallback': 'sc_config',
                    'query_iwt_handling_office': 'xc_config'}
        seen = {}
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in expected:
                seen[node.func.id] = node.args[0].id
        self.assertEqual(seen, expected)

    def test_iwt_helpers_use_passed_connection(self):
        config = {'host': 'WMS'}
        for query in (m.query_iwt_meta, m.query_iwt_to):
            with patch.object(m, 'fetch_all', return_value=[]) as fetch:
                query(config, ['IWT_TEST'])
                self.assertIs(fetch.call_args.args[0], config)
                self.assertIn('t_iwt_head', fetch.call_args.args[1])

    def test_cross_database_join(self):
        from decimal import Decimal
        wms, xc = {'host': 'WMS'}, {'host': 'XC'}
        relationships = [{'iwt_no': 'IWT1', 'house_no': 'H1'},
                         {'iwt_no': 'IWT1', 'house_no': 'H2'},
                         {'iwt_no': 'IWT2', 'house_no': 'H1'}]
        weights = [{'bussiness_no': 'H1', 'chargeable_weight': 100},
                   {'bussiness_no': 'H2', 'chargeable_weight': 300}]
        validation = m.ValidationLog()
        with patch.object(m, 'fetch_all', side_effect=[relationships, weights]) as fetch:
            rows = m.query_iwt_calc(wms, xc, ['IWT1', 'IWT2'], validation)
        self.assertIs(fetch.call_args_list[0].args[0], wms)
        self.assertIs(fetch.call_args_list[1].args[0], xc)
        self.assertNotIn('t_shipment_total', fetch.call_args_list[0].args[1])
        self.assertEqual(fetch.call_args_list[1].args[2], ['H1', 'H2'])
        self.assertEqual([r.do_cw for r in rows], [Decimal(400), Decimal(400), Decimal(100)])
        self.assertFalse(validation.items)

    def test_missing_invalid_duplicate_weights(self):
        relationships = [{'iwt_no': 'IWT1', 'house_no': h} for h in ['H1', 'H2', 'H3', 'H3']]
        weights = [{'bussiness_no': 'H2', 'chargeable_weight': 0},
                   {'bussiness_no': 'H3', 'chargeable_weight': 10},
                   {'bussiness_no': 'H3', 'chargeable_weight': 20}]
        validation = m.ValidationLog()
        with patch.object(m, 'fetch_all', side_effect=[relationships, weights]):
            rows = m.query_iwt_calc({}, {}, ['IWT1'], validation)
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r.h_cw is None for r in rows))
        self.assertEqual({v.issue_type for v in validation.items},
                         {'IWT_CW_NOT_FOUND', 'IWT_CW_INVALID', 'IWT_CW_DUPLICATE', 'IWT_DUPLICATE_HBL'})

    def test_empty_iwt_skips_queries(self):
        with patch.object(m, 'fetch_all') as fetch:
            self.assertEqual(m.query_iwt_calc({}, {}, [], m.ValidationLog()), [])
            fetch.assert_not_called()

    def test_languages(self):
        with patch.object(language, 'LANGUAGE', 'zh'):
            self.assertIn('查询失败', language.translate('WMS database query failed'))
        with patch.object(language, 'LANGUAGE', 'en'):
            self.assertEqual(language.translate('WMS database query failed'), 'WMS database query failed')

    def test_iwt_handling_office_mapping(self):
        records = [{'house_no': h, 'dest_handing_office': o} for h, o in
                   [('H1', None), ('H2', ' Apex-ord '), ('H3', 'APEX-DFW'),
                    ('H3', 'apex-dfw'), ('H4', 'APEX-SFO'), ('H4', 'APEX-LAX')]]
        v = m.ValidationLog()
        config = {'host': 'XC'}
        with patch.object(m, 'fetch_all', return_value=records) as fetch:
            result = m.query_iwt_handling_office(config, ['H1', 'H2', 'H3', 'H3', 'H4'], v)
        self.assertEqual(result, {'h3': 'APEX-DFW', 'h4': 'APEX-LAX; APEX-SFO'})
        self.assertIs(fetch.call_args.args[0], config)
        self.assertEqual(fetch.call_args.args[2], ['H1', 'H2', 'H3', 'H4'])
        self.assertIn('head.dest_handing_office', fetch.call_args.args[1])
        self.assertEqual([item.issue_type for item in v.items], ['IWT_HANDLING_OFFICE_MULTIPLE'])
        with patch.object(m, 'fetch_all') as fetch:
            self.assertEqual(m.query_iwt_handling_office(config, [], v), {})
            fetch.assert_not_called()

    def test_iwt_export_office_does_not_change_billing(self):
        from copy import deepcopy
        from openpyxl import load_workbook
        from test_ord_do_allocation import raw, calc, run
        sources = [raw('DO1', total='100'), raw('IWT1', total='100', pro='P2')]
        shipments = [calc('DO1', '10', customer='CORNING INCORPORATED'), calc('IWT1', '20')]
        shipments[1].hbl_no = 'IWT_HBL'
        v = run(sources, shipments)
        with tempfile.TemporaryDirectory() as folder:
            paths = [Path(folder) / 'before.xlsx', Path(folder) / 'after.xlsx']
            for path, office in zip(paths, [{}, {'iwt_hbl': 'APEX-DFW'}]):
                m.write_workbook(path, sources, shipments, {}, m.CarrierLookup({}, {}), {}, deepcopy(v),
                                 iwt_handling_office=office)
            before, after = [load_workbook(p) for p in paths]
            self.assertEqual(after['IWT']['M1'].value, 'Handling office')
            self.assertEqual(after['IWT']['M2'].value, 'APEX-DFW')
            self.assertEqual(after['IWT'].max_column, 13)
            for name in before.sheetnames:
                first = list(before[name].values)
                second = list(after[name].values)
                if name == 'IWT':
                    first = [r[:12] for r in first]
                    second = [r[:12] for r in second]
                self.assertEqual(first, second, name)
            before.close()
            after.close()


if __name__ == '__main__':
    unittest.main()

