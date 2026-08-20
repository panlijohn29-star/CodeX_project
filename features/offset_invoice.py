import datetime as d
from decimal import Decimal
import os
import re
import uuid

from openpyxl import load_workbook
import pymysql

from platform_config import BASE_DIR, get_db_config


ALLOWED_DB_PROFILES = ("scdbus", "scdbca")
OUTPUT_DIR = os.path.join(BASE_DIR, "reports_v4", "offset_invoice")
TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "offset_inv_temp.xlsx")
MAX_INVOICES_PER_QUERY = 500

SHEET_HEADERS = [
    "AP# (Credit Note#)", "OFFICE", "Account(billing party)",
    "Vendor Invoice Date", "Vendor Invoice Number", "JOB NO", "RP TYPE",
    "Charge display name", "Charge Category ", "Charge Code", "Amount", "UserName",
]

FEATURE = {
    "id": "offset_invoice",
    "title": "Offset Invoice",
    "category": "Tools",
    "description": "Generate AR/AP offset invoice upload workbook from invoice numbers.",
    "supports_cancel": False,
    "output_type": "interactive",
    "input_schema": [],
    "template": "offset_invoice.html",
    "db_profiles": ALLOWED_DB_PROFILES,
}

AR_SQL = """
SELECT c.INVOICE_NO AS source_invoice_no, c.company_code, c.balance,
       c.customer_invoice_no, v.job_no AS job_no, charge_local_name, charge_category,
       charge_type, exchange_usd * -1 AS amount
FROM CF_CHARGES c
LEFT JOIN V_JOBINFO v ON v.JOB_ID = c.JOB_ID
WHERE c.INVOICE_NO IN ({placeholders})
"""

AP_SQL = """
SELECT c.INVOICE_NO AS source_invoice_no, c.company_code, c.balance,
       c.customer_invoice_no, v.job_no AS job_no, charge_local_name, charge_category,
       charge_type, exchange_usd * -1 AS amount
FROM CF_COST c
LEFT JOIN V_JOBINFO v ON v.JOB_ID = c.JOB_ID
WHERE c.INVOICE_NO IN ({placeholders})
"""


def normalize_db_profile(db_profile):
    profile = str(db_profile or "").strip().lower()
    if profile not in ALLOWED_DB_PROFILES:
        raise ValueError("Please select scdbus or scdbca")
    return profile


def normalize_username(username):
    value = str(username or "").strip()
    if not value:
        raise ValueError("UserName is required")
    return value


def normalize_invoices(value):
    raw_values = [value] if isinstance(value, str) else list(value or [])
    invoices = []
    seen = set()
    for raw in raw_values:
        for item in re.split(r"[\r\n,;\t]+", str(raw or "")):
            invoice = item.strip()
            key = invoice.upper()
            if invoice and key not in seen:
                invoices.append(invoice)
                seen.add(key)
    if not invoices:
        raise ValueError("Please enter at least one INV#")
    return invoices


def _connect(db_profile):
    config = get_db_config(db_profile)
    config["cursorclass"] = pymysql.cursors.DictCursor
    return pymysql.connect(**config)


def query_rows(db_profile, invoices, sql):
    rows = []
    connection = _connect(db_profile)
    try:
        with connection.cursor() as cursor:
            for start in range(0, len(invoices), MAX_INVOICES_PER_QUERY):
                batch = invoices[start:start + MAX_INVOICES_PER_QUERY]
                placeholders = ", ".join(["%s"] * len(batch))
                cursor.execute(sql.format(placeholders=placeholders), batch)
                rows.extend(cursor.fetchall())
    finally:
        connection.close()
    return rows


def as_amount(value):
    if value is None:
        return None
    return float(Decimal(str(value)))


def output_values(rows, rp_type, username):
    return [[
        None,
        row.get("company_code"),
        row.get("balance"),
        None,
        row.get("customer_invoice_no"),
        row.get("job_no"),
        rp_type,
        row.get("charge_local_name"),
        row.get("charge_category"),
        row.get("charge_type"),
        as_amount(row.get("amount")),
        username,
    ] for row in rows]


def create_workbook(output_path, ar_rows, ap_rows, username):
    if not os.path.exists(TEMPLATE_PATH):
        raise ValueError("Offset invoice template was not found")
    workbook = load_workbook(TEMPLATE_PATH)
    try:
        for sheet_name, rows, rp_type in (("AR", ar_rows, "AR"), ("AP", ap_rows, "AP")):
            if sheet_name not in workbook.sheetnames:
                raise ValueError("Template is missing {0} sheet".format(sheet_name))
            sheet = workbook[sheet_name]
            actual_headers = [sheet.cell(1, column).value for column in range(1, 13)]
            if actual_headers != SHEET_HEADERS:
                raise ValueError("Template {0} headers do not match the expected format".format(sheet_name))
            if sheet.max_row > 1:
                sheet.delete_rows(2, sheet.max_row - 1)
            for row in output_values(rows, rp_type, username):
                sheet.append(row)
                current = sheet.max_row
                for column in (1, 5, 6, 7, 10, 12):
                    sheet.cell(current, column).number_format = "@"
                sheet.cell(current, 11).number_format = "#,##0.00"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = "A1:L{0}".format(max(sheet.max_row, 1))
        workbook.save(output_path)
    finally:
        workbook.close()


def _json_safe_rows(rows):
    return [
        {str(key): "" if value is None else str(value) for key, value in row.items()}
        for row in rows
    ]


def generate_payload(db_profile, username, invoice_text):
    profile = normalize_db_profile(db_profile)
    user = normalize_username(username)
    invoices = normalize_invoices(invoice_text)

    ar_rows = query_rows(profile, invoices, AR_SQL)
    ap_rows = query_rows(profile, invoices, AP_SQL)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_name = "offset_invoice_{0}_{1}.xlsx".format(
        d.datetime.now().strftime("%Y%m%d_%H%M%S"),
        uuid.uuid4().hex[:8],
    )
    output_path = os.path.join(OUTPUT_DIR, output_name)
    create_workbook(output_path, ar_rows, ap_rows, user)

    matched = {
        str(row.get("source_invoice_no") or "").strip().upper()
        for row in [*ar_rows, *ap_rows]
    }
    unmatched = [invoice for invoice in invoices if invoice.upper() not in matched]
    ar_total = sum((as_amount(row.get("amount")) or 0) for row in ar_rows)
    ap_total = sum((as_amount(row.get("amount")) or 0) for row in ap_rows)

    return {
        "ok": True,
        "db_profile": profile,
        "username": user,
        "invoice_count": len(invoices),
        "ar_count": len(ar_rows),
        "ap_count": len(ap_rows),
        "ar_total": "{0:,.2f}".format(ar_total),
        "ap_total": "{0:,.2f}".format(ap_total),
        "unmatched": unmatched,
        "download_name": output_name,
        "download_url": "/download/offset-invoice/{0}".format(output_name),
        "ar_preview_rows": _json_safe_rows(ar_rows[:25]),
        "ap_preview_rows": _json_safe_rows(ap_rows[:25]),
    }


def output_path_for(filename):
    safe_name = os.path.basename(str(filename or ""))
    if not safe_name or safe_name != filename:
        return None
    path = os.path.abspath(os.path.join(OUTPUT_DIR, safe_name))
    output_root = os.path.abspath(OUTPUT_DIR)
    if not path.startswith(output_root + os.sep):
        return None
    return path
