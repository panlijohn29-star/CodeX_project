import datetime as d
from io import BytesIO
from pathlib import Path

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from platform_config import get_db_config


ALLOWED_DB_PROFILES = ("scdbus",)
SQL_PATH = Path(__file__).with_name("ord_ae_closing.sql")
DEFAULT_COLUMNS = [
    "JOB_NO", "CUSTOMER", "MBL_NO", "HBL_NO", "ETD", "CLOSING_STATUS",
    "SERVICE_TYPE", "PAYMENT_TERMS", "AR_BILL", "AP_CNS", "AP_TRUCK", "AP_ATC",
    "AP_WH", "ORD_ARDONE", "ORD_ARCHECK", "ORD_APDONE", "ORD_APCHECK",
    "Total_charge", "Total_cost", "Total_GP",
]
COLUMN_LABELS = {
    "JOB_NO": "JOB#", "MBL_NO": "MBL", "HBL_NO": "HBL", "CLOSING_STATUS": "STATUS",
    "SERVICE_TYPE": "SVC", "PAYMENT_TERMS": "PAY", "AR_BILL": "AR", "AP_CNS": "CNS",
    "AP_TRUCK": "TRK", "AP_ATC": "ATC", "AP_WH": "WH", "ORD_ARDONE": "AR DONE",
    "ORD_ARCHECK": "AR CHK", "ORD_APDONE": "AP DONE", "ORD_APCHECK": "AP CHK",
    "Total_charge": "CHG", "Total_cost": "COST", "Total_GP": "GP",
}

FEATURE = {
    "id": "ord_ae_closing",
    "title": "ORD AE Closing Report",
    "category": "Reports",
    "description": "Filter ORD air-export jobs, review closing status and GP, then export the result or batch-closing file.",
    "supports_cancel": False,
    "output_type": "interactive",
    "input_schema": [],
    "template": "ord_ae_closing.html",
    "db_profiles": ALLOWED_DB_PROFILES,
}


def normalize_db_profile(value):
    profile = str(value or "").strip().lower()
    if profile not in ALLOWED_DB_PROFILES:
        raise ValueError("Please select scdbus")
    return profile


def normalize_date(value, label, required=False):
    text = str(value or "").strip()
    if not text:
        if required:
            raise ValueError("Please fill in {0} before running the report.".format(label))
        return ""
    try:
        return d.date.fromisoformat(text).isoformat()
    except ValueError:
        raise ValueError("{0} must use YYYY-MM-DD format".format(label))


def normalize_status(value):
    status = str(value or "").strip().upper()
    if status not in ("", "CLOSED", "UNCLOSED"):
        raise ValueError("Please select All, CLOSED, or UNCLOSED")
    return status


def _like(value, mode="contains"):
    text = str(value or "").strip()
    if not text:
        return ""
    return "{0}%".format(text) if mode == "starts_with" else "%{0}%".format(text)


def build_query(closing_status="", etd_from="", etd_to="", job_no="", customer="", customer_mode="starts_with", mbl_no="", hbl_no=""):
    status = normalize_status(closing_status)
    start_date = normalize_date(etd_from, "ETD From", required=True)
    end_date = normalize_date(etd_to, "ETD To")
    clauses, params = ["AND DATE(v_jobinfo.ETD) >= %s"], [start_date]
    if end_date:
        clauses.append("AND DATE(v_jobinfo.ETD) <= %s")
        params.append(end_date)
    if status == "CLOSED":
        clauses.extend(["AND COALESCE(OCB.B_CHECK_CHARGES, 0) = 1", "AND COALESCE(OCB.B_CHECK_COST, 0) = 1"])
    elif status == "UNCLOSED":
        clauses.append("AND (COALESCE(OCB.B_CHECK_CHARGES, 0) <> 1 OR COALESCE(OCB.B_CHECK_COST, 0) <> 1)")
    for column, value, mode in (("v_jobinfo.JOB_NO", job_no, "contains"), ("v_jobinfo.CUSTOMER", customer, customer_mode), ("v_jobinfo.MBL_NO", mbl_no, "contains"), ("v_jobinfo.HBL_NO", hbl_no, "contains")):
        match = _like(value, "starts_with" if mode == "starts_with" else "contains")
        if match:
            clauses.append("AND {0} LIKE %s".format(column))
            params.append(match)
    sql = SQL_PATH.read_text(encoding="utf-8").replace("<CONDITION>", "\n".join(clauses) + "\nORDER BY ETD DESC")
    return sql, params, status


def _connect(profile):
    config = get_db_config(profile)
    config["cursorclass"] = pymysql.cursors.DictCursor
    return pymysql.connect(**config)


def _literal(value):
    return "'{0}'".format(str(value).replace("\\", "\\\\").replace("'", "''"))


def preview_query(sql, params):
    preview = sql.replace("%%", "%")
    for value in params:
        preview = preview.replace("%s", _literal(value), 1)
    return preview


def _public_row(row):
    return {str(key): "" if value is None else str(value) for key, value in row.items()}


def _is_closed(row):
    return int(row.get("ORD_ARCHECK") or 0) == 1 and int(row.get("ORD_APCHECK") or 0) == 1


def _as_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _with_status(rows):
    result = []
    for source in rows:
        row = dict(source)
        row["CLOSING_STATUS"] = "CLOSED" if _is_closed(row) else "UNCLOSED"
        result.append(row)
    return result


def summarize(rows):
    closed_rows = [row for row in rows if _is_closed(row)]
    unclosed_rows = [row for row in rows if not _is_closed(row)]
    return {
        "total": len(rows), "closed": len(closed_rows), "unclosed": len(unclosed_rows),
        "total_gp": sum(_as_float(row.get("Total_GP")) for row in rows),
        "closed_gp": sum(_as_float(row.get("Total_GP")) for row in closed_rows),
        "unclosed_gp": sum(_as_float(row.get("Total_GP")) for row in unclosed_rows),
    }


def _run(profile, **filters):
    sql, params, status = build_query(**filters)
    connection = _connect(profile)
    try:
        with connection.cursor() as cursor:
            cursor.execute(sql.replace("%", "%%").replace("%%s", "%s"), params)
            rows = _with_status(cursor.fetchall())
    finally:
        connection.close()
    return rows, sql, params, status


def search_payload(db_profile, **filters):
    profile = normalize_db_profile(db_profile)
    rows, sql, params, status = _run(profile, **filters)
    public_rows = [_public_row(row) for row in rows]
    columns = list(public_rows[0]) if public_rows else list(DEFAULT_COLUMNS)
    if "CLOSING_STATUS" in columns:
        columns.remove("CLOSING_STATUS")
    insert_at = columns.index("ETD") + 1 if "ETD" in columns else len(columns)
    columns.insert(insert_at, "CLOSING_STATUS")
    return {"ok": True, "db_profile": profile, "closing_status": status, "columns": columns, "column_labels": COLUMN_LABELS, "rows": public_rows, "row_count": len(public_rows), "summary": summarize(rows), "query_sql": preview_query(sql, params)}


def preview_payload(db_profile, **filters):
    profile = normalize_db_profile(db_profile)
    sql, params, status = build_query(**filters)
    return {"ok": True, "db_profile": profile, "closing_status": status, "query_sql": preview_query(sql, params)}


def _style(sheet):
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor="0F766E")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 32)


def export_workbook(db_profile, export_type, request_date="", **filters):
    profile = normalize_db_profile(db_profile)
    rows, _, _, _ = _run(profile, **filters)
    workbook = Workbook()
    sheet = workbook.active
    if export_type == "batch_closing":
        sheet.title = "Batch Closing"
        sheet.append(["JOB#", "Close Office", "REQUEST DATE"])
        date_value = normalize_date(request_date, "Request date") or d.date.today().isoformat()
        for row in rows:
            sheet.append([row.get("JOB_NO", ""), "ORD", date_value])
        filename = "ORDAE_batch_closing_{0}.xlsx".format(d.datetime.now().strftime("%Y%m%d_%H%M%S"))
    elif export_type == "original":
        sheet.title = "Report Output"
        sheet.append([COLUMN_LABELS.get(column, column) for column in DEFAULT_COLUMNS])
        for row in rows:
            sheet.append([row.get(column, "") for column in DEFAULT_COLUMNS])
        filename = "ORDAE_original_{0}.xlsx".format(d.datetime.now().strftime("%Y%m%d_%H%M%S"))
    else:
        raise ValueError("Unknown export type")
    _style(sheet)
    output = BytesIO(); workbook.save(output); output.seek(0)
    return output, filename
