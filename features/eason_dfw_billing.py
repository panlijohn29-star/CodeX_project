"""Eason DFW billing workbook processor and converter."""
from __future__ import annotations

import io
import json
import os
import re
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pymysql
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from platform_config import BASE_DIR, get_db_config


SOURCE_SHEET = "Sheet1"
RULE_SHEET = "Billing Sheet"
PROCESSED_SHEET = "Processed Billing"
UNMATCHED_SHEET = "Unmatched Customers"
ASSET_DIR = Path(__file__).with_name("eason_dfw_billing_assets")
ACCOUNT_MAP_PATH = ASSET_DIR / "billing_party_mapping.json"
CHARGE_MAP_PATH = ASSET_DIR / "charge_code_mapping.json"
TEMPLATE_PATH = ASSET_DIR / "conversion_target.xlsx"
OUTPUT_DIR = Path(BASE_DIR) / "reports_v4" / "eason_dfw_billing"
INITIAL_CUSTOMER_LOG_PATH = ASSET_DIR / "customer_import_log.json"
CUSTOMER_LOG_PATH = OUTPUT_DIR / "customer_import_log.json"
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
DB_FIELDS = ["DB_THC", "PACKAGE_NUM", "PACKAGE_TYPE", "CW"]
PLT_PACKAGE_TYPES = {"PLTS", "PLT", "PALLET", "PALLETS"}
NON_CHARGE_COLUMNS = {"CUSTOMER", "MONTH", "PERIOD", "ETD_ETA", "ETD", "ETA", "JOB_NO", "MBL_NO", "HBL_NO", "DB_THC", "PACKAGE_NUM", "PACKAGE_TYPE", "CW", "EDI INVOICE", "REGION", "BILLING PARTY", "MATCH STATUS"}
TARGET_HEADERS = ["Customer", "Charge Code", "HouseNo/BOLNo/InternalNo", "JobType", "Account", "Invoice Title", "Invoice Title Name", "Currency", "Unit", "Quantity", "Unit Price", "RPType", "Remark", "Vendor invoice#", "Vendor invoice Date", "CHARGE_CATEGORY", "CHARGE_LOCAL_NAME"]

FEATURE = {"id": "eason_dfw_billing", "title": "Eason_DFW_Billing", "category": "Tools", "description": "Process a DFW billing workbook and generate full, DO, and NODO upload files.", "supports_cancel": False, "output_type": "interactive", "input_schema": [], "template": "eason_dfw_billing.html", "db_profiles": ("scdbus",)}


class AuditError(ValueError):
    def __init__(self, message, issues):
        super().__init__(message)
        self.issues = issues


@dataclass
class JobChargeData:
    thc: Decimal = Decimal("0")
    package_num: object = None
    package_type: object = None
    cw: object = None


def _load_json(path):
    if not path.exists():
        raise ValueError("Required Eason DFW Billing configuration is missing: {0}".format(path.name))
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("Invalid configuration: {0}".format(path.name))
    return value


def _save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _parse_account_mappings(raw):
    mappings = {}
    for line in str(raw or "").splitlines():
        if not line.strip():
            continue
        if "=" not in line:
            raise ValueError("Account mappings must use CODE=Full Account Name")
        code, account = (part.strip() for part in line.split("=", 1))
        if not code or not account:
            raise ValueError("Account mappings must include both code and account name")
        mappings[code] = account
    return mappings


def _parse_charge_mappings(raw):
    mappings = {}
    for line in str(raw or "").splitlines():
        if not line.strip():
            continue
        parts = [part.strip() for part in line.split("|")]
        if len(parts) != 4 or not all(parts):
            raise ValueError("Charge mappings must use COLUMN|Charge Code|Category|Local Name")
        try:
            charge_code = int(parts[1])
        except ValueError as exc:
            raise ValueError("Charge Code must be a number") from exc
        mappings[parts[0].upper()] = {"Charge Code": charge_code, "CHARGE_CATEGORY": parts[2], "CHARGE_LOCAL_NAME": parts[3]}
    return mappings


def mappings_payload():
    accounts = _load_json(ACCOUNT_MAP_PATH)
    charges = _load_json(CHARGE_MAP_PATH)
    return {"accounts_text": "\n".join("{0}={1}".format(code, account) for code, account in sorted(accounts.items())), "charges_text": "\n".join("{0}|{1}|{2}|{3}".format(column, value.get("Charge Code", ""), value.get("CHARGE_CATEGORY", ""), value.get("CHARGE_LOCAL_NAME", "")) for column, value in sorted(charges.items()))}


def _load_customer_import_log():
    if CUSTOMER_LOG_PATH.exists():
        return _load_json(CUSTOMER_LOG_PATH)
    if INITIAL_CUSTOMER_LOG_PATH.exists():
        return _load_json(INITIAL_CUSTOMER_LOG_PATH)
    return {"known_customers": [], "runs": []}


def _update_customer_import_log(source_name, customers):
    log_data = _load_customer_import_log()
    known_customers = set(log_data.get("known_customers", []))
    imported_customers = sorted({customer for customer in customers if customer})
    new_customers = [customer for customer in imported_customers if customer not in known_customers]
    log_data["known_customers"] = sorted(known_customers.union(imported_customers))
    runs = list(log_data.get("runs", []))
    runs.append({"timestamp": datetime.now().isoformat(timespec="seconds"), "source_file": source_name, "customer_count": len(imported_customers), "customers": imported_customers, "new_customers": new_customers})
    log_data["runs"] = runs
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _save_json(CUSTOMER_LOG_PATH, log_data)
    return imported_customers, new_customers


def _decimal(value):
    return Decimal("0") if value in (None, "") else Decimal(str(value))


def _number(value):
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral() else float(value)
    return value


def _numeric_text(value):
    if not isinstance(value, str):
        return _number(value)
    text = value.strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)
    return value


def _headers(sheet):
    return {str(cell.value).strip(): index for index, cell in enumerate(sheet[1]) if cell.value is not None}


def _value(row, index, name, default=None):
    position = index.get(name)
    return default if position is None or position >= len(row) else row[position]


def _rule_map(sheet):
    headers = [cell.value for cell in sheet[1]]
    rules = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            rules[str(row[0]).strip()] = {headers[i]: row[i] for i in range(1, len(headers))}
    return headers, rules


def _calculate_atc(rule, package_type, cw):
    if rule in (None, "") or cw in (None, ""):
        return _numeric_text(rule)
    text = str(rule).upper().strip()
    plt = re.search(r"(\d+(?:\.\d+)?)\s*PLT", text)
    loose = re.search(r"(\d+(?:\.\d+)?)\s*LOOSE", text)
    minimum = re.search(r"MIN\s*(\d+(?:\.\d+)?)", text)
    if not plt and not loose:
        return _numeric_text(rule)
    package = str(package_type or "").upper().strip()
    rate = Decimal(plt.group(1)) if package in PLT_PACKAGE_TYPES and plt else Decimal(loose.group(1)) if loose else None
    if rate is None:
        return _numeric_text(rule)
    amount = rate * _decimal(cw)
    if minimum:
        amount = max(amount, Decimal(minimum.group(1)))
    return _number(amount.quantize(Decimal("0.01")))


def _handling(rule, customer, mawb, seen):
    if rule is None:
        return None
    text = str(rule).strip()
    if "/" not in text:
        return _numeric_text(rule)
    amount, basis = (part.strip() for part in text.split("/", 1))
    if basis.upper() == "HAWB":
        return _numeric_text(amount)
    if basis.upper() == "MAWB":
        key = (customer, mawb, amount)
        if key in seen:
            return ""
        seen.add(key)
        return _numeric_text(amount)
    return _numeric_text(rule)


def _fetch_job_data(job_numbers):
    if not job_numbers:
        return {}
    unique_jobs = list(dict.fromkeys(job_numbers))
    placeholders = ", ".join(["%s"] * len(unique_jobs))
    package_sql = "SELECT JOB_NO, HAWB_PKGS_NUM, HAWB_PKGS_TYPE, HAWB_CHARGEABLE_WEIGHT AS CW FROM V_JOBINFO WHERE JOB_NO IN ({0})".format(placeholders)
    thc_sql = "SELECT v_jobinfo.JOB_NO, cf_cost.EXCHANGE_USD AS THC FROM CF_COST LEFT JOIN V_JOBINFO ON CF_COST.JOB_ID=V_JOBINFO.JOB_ID WHERE CF_COST.COMPANY_CODE=%s AND cf_cost.CHARGE_LOCAL_NAME LIKE %s AND v_jobinfo.JOB_NO IN ({0})".format(placeholders)
    config = get_db_config("scdbus")
    config["cursorclass"] = pymysql.cursors.SSDictCursor
    result = {}
    connection = pymysql.connect(**config)
    try:
        with connection.cursor() as cursor:
            cursor.execute(package_sql, unique_jobs)
            for row in cursor.fetchall():
                job = row.get("JOB_NO")
                if job:
                    item = result.setdefault(job, JobChargeData())
                    item.package_num, item.package_type, item.cw = row.get("HAWB_PKGS_NUM"), row.get("HAWB_PKGS_TYPE"), row.get("CW")
            cursor.execute(thc_sql, ["APEX-DFW", "terminal%", *unique_jobs])
            for row in cursor.fetchall():
                job = row.get("JOB_NO")
                if job:
                    result.setdefault(job, JobChargeData()).thc += _decimal(row.get("THC"))
    finally:
        connection.close()
    return result


def _autosize(sheet):
    for column in sheet.columns:
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max((len(str(cell.value)) for cell in column if cell.value is not None), default=0) + 2, 28)


def _process(workbook, account_map):
    if SOURCE_SHEET not in workbook.sheetnames or RULE_SHEET not in workbook.sheetnames:
        raise ValueError("Workbook must include Sheet1 and Billing Sheet")
    source, billing = workbook[SOURCE_SHEET], workbook[RULE_SHEET]
    source_index = _headers(source)
    missing_columns = [name for name in ("CUSTOMER", "JOB_NO", "MBL_NO", "HBL_NO") if name not in source_index]
    if missing_columns:
        raise ValueError("Sheet1 is missing required columns: {0}".format(", ".join(missing_columns)))
    rule_headers, rules = _rule_map(billing)
    missing_codes = sorted({str(rule.get("Billing Party")).strip() for rule in rules.values() if rule.get("Billing Party") not in (None, "") and str(rule.get("Billing Party")).strip() not in account_map})
    if missing_codes:
        raise ValueError("Missing Billing Party account mapping for: {0}. Add CODE=Full Account Name and run again.".format(", ".join(missing_codes)))
    rows = list(source.iter_rows(min_row=2, values_only=True))
    imported_customers = [str(_value(row, source_index, "CUSTOMER")).strip() for row in rows if _value(row, source_index, "CUSTOMER") is not None]
    jobs = [str(_value(row, source_index, "JOB_NO")).strip() for row in rows if _value(row, source_index, "JOB_NO") is not None]
    job_data = _fetch_job_data(jobs)
    for name in (PROCESSED_SHEET, UNMATCHED_SHEET):
        if name in workbook.sheetnames:
            del workbook[name]
    processed, unmatched = workbook.create_sheet(PROCESSED_SHEET), workbook.create_sheet(UNMATCHED_SHEET)
    processed.append([cell.value for cell in source[1]] + DB_FIELDS + rule_headers[1:] + ["Match Status"])
    unmatched.append(["CUSTOMER", "JOB_NO"])
    seen = set()
    unmatched_records = []
    for row in rows:
        customer = str(_value(row, source_index, "CUSTOMER") or "").strip()
        job, mawb = str(_value(row, source_index, "JOB_NO") or "").strip(), str(_value(row, source_index, "MBL_NO") or "").strip()
        rule, data = rules.get(customer), job_data.get(job, JobChargeData())
        db_values = [_number(data.thc), data.package_num, data.package_type, data.cw]
        if rule is None:
            unmatched_records.append({"customer": customer or "(blank customer)", "job_no": job or "(blank job)"})
            unmatched.append([customer, job])
            processed.append(list(row) + db_values + [None] * (len(rule_headers) - 1) + ["UNMATCHED"])
            continue
        values = [_numeric_text(rule.get(header)) for header in rule_headers[1:]]
        positions = {header: i for i, header in enumerate(rule_headers[1:])}
        if "THC" in positions: values[positions["THC"]] = _number(data.thc)
        if "Billing Party" in positions: values[positions["Billing Party"]] = account_map[str(rule.get("Billing Party")).strip()]
        if "HANDLING" in positions: values[positions["HANDLING"]] = _handling(rule.get("HANDLING"), customer, mawb, seen)
        if "ATC" in positions: values[positions["ATC"]] = _calculate_atc(rule.get("ATC"), data.package_type, data.cw)
        processed.append(list(row) + db_values + values + ["MATCHED"])
    if not unmatched_records:
        unmatched.append(["All customers matched", ""])
    for sheet in (processed, unmatched):
        for cell in sheet[1]: cell.font = Font(bold=True)
        _autosize(sheet)
    return processed, unmatched_records, imported_customers


def _charge_columns(processed):
    return [str(cell.value).strip() for cell in processed[1] if cell.value is not None and str(cell.value).strip().upper() not in NON_CHARGE_COLUMNS]


def _output_rows(processed, charge_map):
    headers = [cell.value for cell in processed[1]]
    index = {header: i for i, header in enumerate(headers)}
    charge_columns = _charge_columns(processed)
    missing = sorted(column for column in charge_columns if column.upper() not in charge_map)
    if missing:
        raise ValueError("Missing charge mapping for: {0}. Add COLUMN|Charge Code|Category|Local Name and run again.".format(", ".join(missing)))
    rows = []
    for row in processed.iter_rows(min_row=2, values_only=True):
        for column in charge_columns:
            value = _calculate_atc(row[index["ATC"]], row[index["PACKAGE_TYPE"]], row[index["CW"]]) if column == "ATC" else _numeric_text(row[index[column]])
            if value in (None, "", 0): continue
            mapping = charge_map[column.upper()]
            account, hbl = row[index["Billing Party"]], row[index["HBL_NO"]]
            rows.append({"Customer": row[index["CUSTOMER"]], "Charge Code": mapping["Charge Code"], "HouseNo/BOLNo/InternalNo": hbl, "JobType": "AI", "Account": account, "Invoice Title": account, "Invoice Title Name": account, "Currency": "USD", "Unit": "Job", "Quantity": 1, "Unit Price": value, "RPType": "AR", "Remark": hbl, "Vendor invoice#": None, "Vendor invoice Date": None, "CHARGE_CATEGORY": mapping["CHARGE_CATEGORY"], "CHARGE_LOCAL_NAME": mapping["CHARGE_LOCAL_NAME"]})
    return rows


def _write_conversion(rows, path, sheet_name):
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook[workbook.sheetnames[0]]
    sheet.title = sheet_name
    for row in range(2, sheet.max_row + 1):
        for column in range(1, len(TARGET_HEADERS) + 1):
            sheet.cell(row, column).value = None
    if sheet.max_row > 2:
        sheet.delete_rows(3, sheet.max_row - 2)
    for row_number, record in enumerate(rows, start=2):
        for column in range(1, len(TARGET_HEADERS) + 1):
            source = sheet.cell(2, column)
            target = sheet.cell(row_number, column)
            if row_number != 2:
                target._style = copy(source._style)
                target.number_format = source.number_format
                target.font = copy(source.font)
                target.fill = copy(source.fill)
                target.border = copy(source.border)
                target.alignment = copy(source.alignment)
                target.protection = copy(source.protection)
            target.value = record.get(TARGET_HEADERS[column - 1])
    workbook.save(path)
    workbook.close()


def _audit_rows(rows, do_rows, nodo_rows):
    issues = []
    if not rows:
        issues.append("No billable charge rows were generated.")
    if len(rows) != len(do_rows) + len(nodo_rows):
        issues.append("Full, DO, and NODO row counts do not reconcile.")
    do_hbls = {row["HouseNo/BOLNo/InternalNo"] for row in do_rows}
    nodo_hbls = {row["HouseNo/BOLNo/InternalNo"] for row in nodo_rows}
    if do_hbls.intersection(nodo_hbls):
        issues.append("A HAWB appears in both DO and NODO output.")
    delivery_hbls = {row["HouseNo/BOLNo/InternalNo"] for row in rows if row["CHARGE_LOCAL_NAME"] == "Delivery"}
    if do_hbls != delivery_hbls:
        issues.append("DO output does not match HAWBs that contain a Delivery charge.")
    for index, row in enumerate(rows, start=2):
        missing = [field for field in ("Customer", "Charge Code", "HouseNo/BOLNo/InternalNo", "Account", "CHARGE_CATEGORY", "CHARGE_LOCAL_NAME") if row.get(field) in (None, "")]
        if missing:
            issues.append("Converted row {0} is missing: {1}.".format(index, ", ".join(missing)))
            continue
        if re.fullmatch(r"A\d+", str(row["Account"]).strip(), flags=re.IGNORECASE):
            issues.append("Converted row {0} has an unmapped A-code account: {1}.".format(index, row["Account"]))
        try:
            if int(row["Charge Code"]) <= 0:
                issues.append("Converted row {0} has an invalid Charge Code.".format(index))
        except (TypeError, ValueError):
            issues.append("Converted row {0} has an invalid Charge Code.".format(index))
    return issues


def _audit_output_workbook(path, expected_sheet, expected_rows):
    issues = []
    if not path.exists() or not path.is_file() or path.stat().st_size == 0:
        return ["Output file was not created: {0}.".format(path.name)]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [expected_sheet]:
            issues.append("{0} has an unexpected worksheet layout.".format(path.name))
            return issues
        sheet = workbook[expected_sheet]
        headers = [cell.value for cell in sheet[1]]
        if headers != TARGET_HEADERS:
            issues.append("{0} headers do not match the billing upload template.".format(path.name))
        actual_rows = sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row))
        if actual_rows != expected_rows:
            issues.append("{0} contains {1} data rows; expected {2}.".format(path.name, actual_rows, expected_rows))
    finally:
        workbook.close()
    return issues


def _remove_outputs(paths):
    for path in paths:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            pass


def generate_payload(uploaded, account_text, charge_text):
    if not uploaded or not uploaded.filename:
        raise ValueError("Select a source .xlsx workbook")
    if not uploaded.filename.lower().endswith(".xlsx") or uploaded.filename.startswith("~$"):
        raise ValueError("Only a normal .xlsx workbook can be processed")
    content = uploaded.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES: raise ValueError("Workbook exceeds the 25 MB limit")
    account_map, charge_map = _load_json(ACCOUNT_MAP_PATH), _load_json(CHARGE_MAP_PATH)
    account_map.update(_parse_account_mappings(account_text)); charge_map.update(_parse_charge_mappings(charge_text))
    try:
        workbook = load_workbook(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("Unable to open the uploaded workbook") from exc
    processed, unmatched_records, imported_customers = _process(workbook, account_map)
    if unmatched_records:
        unique_unmatched = []
        seen_unmatched = set()
        for item in unmatched_records:
            key = (item["customer"], item["job_no"])
            if key not in seen_unmatched:
                unique_unmatched.append(item)
                seen_unmatched.add(key)
        raise AuditError(
            "Audit failed: {0} unmatched customer record(s). Add the customer to Billing Sheet, then run again.".format(len(unique_unmatched)),
            [{"type": "unmatched_customer", "customer": item["customer"], "job_no": item["job_no"]} for item in unique_unmatched],
        )
    rows = _output_rows(processed, charge_map)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    token = uuid.uuid4().hex[:10]
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(uploaded.filename).stem).strip("._") or "billing"
    processed_path = OUTPUT_DIR / "{0}_{1}_processed.xlsx".format(stem, token)
    do_rows = [row for row in rows if any(other["HouseNo/BOLNo/InternalNo"] == row["HouseNo/BOLNo/InternalNo"] and other["CHARGE_LOCAL_NAME"] == "Delivery" for other in rows)]
    nodo_rows = [row for row in rows if row not in do_rows]
    full_path, do_path, nodo_path = (OUTPUT_DIR / "{0}_{1}.xlsx".format(stem, token), OUTPUT_DIR / "{0}_{1}_DO.xlsx".format(stem, token), OUTPUT_DIR / "{0}_{1}_NODO.xlsx".format(stem, token))
    output_paths = (processed_path, full_path, do_path, nodo_path)
    row_issues = _audit_rows(rows, do_rows, nodo_rows)
    if row_issues:
        raise AuditError("Audit failed. No Excel files were created.", [{"type": "data_check", "message": issue} for issue in row_issues])
    try:
        workbook.save(processed_path)
        _write_conversion(rows, full_path, "converted sheet")
        _write_conversion(do_rows, do_path, "DO")
        _write_conversion(nodo_rows, nodo_path, "NODO")
        file_issues = []
        for path, sheet_name, expected_rows in ((full_path, "converted sheet", len(rows)), (do_path, "DO", len(do_rows)), (nodo_path, "NODO", len(nodo_rows))):
            file_issues.extend(_audit_output_workbook(path, sheet_name, expected_rows))
        if file_issues:
            raise AuditError("Audit failed. Excel downloads were blocked.", [{"type": "file_check", "message": issue} for issue in file_issues])
    except Exception:
        _remove_outputs(output_paths)
        raise
    _save_json(ACCOUNT_MAP_PATH, dict(sorted(account_map.items()))); _save_json(CHARGE_MAP_PATH, dict(sorted(charge_map.items())))
    imported_customers, new_customers = _update_customer_import_log(uploaded.filename, imported_customers)
    return {"ok": True, "source_name": uploaded.filename, "processed_rows": processed.max_row - 1, "unmatched_count": 0, "imported_customer_count": len(imported_customers), "new_customers": new_customers, "converted_rows": len(rows), "do_rows": len(do_rows), "nodo_rows": len(nodo_rows), "outputs": [{"name": path.name, "download_url": "/download/eason-dfw-billing/{0}".format(path.name)} for path in output_paths]}


def output_path_for(filename):
    safe = os.path.basename(str(filename or ""))
    path = (OUTPUT_DIR / safe).resolve()
    return path if safe == filename and path.parent == OUTPUT_DIR.resolve() else None
