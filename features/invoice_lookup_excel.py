"""Excel snapshot export. All business decisions belong to lookup_core."""
from __future__ import annotations

import hashlib
from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from features.invoice_lookup_core import LOOKUP_HEADERS, RAW_HEADERS, LookupError, Result, key

NAVY = "203A52"
RED = "FCE4E4"
AMBER = "FFF0CC"
GRAY = "EDF0F3"
STATUS_LABELS = {"PASS": "PASS", "FAIL": "FAIL", "INCOMPLETE": "未完成", "UNDETERMINED": "无法确定", "EXCLUDED_AR": "AR 正常排除"}


def put(ws, row, column, value):
    if isinstance(value, Decimal):
        value = float(value)
    cell = ws.cell(row, column)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"  # Imported identifiers/remarks must never become executable formulas.
    else:
        cell.value = value
    cell.font = Font(name="Arial", size=10, color="202B33")
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    if isinstance(value, (int, float)):
        cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    return cell


def append(ws, values):
    row = ws.max_row + 1 if ws.cell(1, 1).value is not None else 1
    for col, value in enumerate(values, 1):
        put(ws, row, col, value)
    return row


def heading(ws, row, headers):
    for col, value in enumerate(headers, 1):
        cell = put(ws, row, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def table(ws, header_row, end_row, columns, name):
    if end_row > header_row:
        item = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(columns)}{end_row}")
        item.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(item)


def status(issues):
    return "FAIL" if any(i.severity == "FAIL" for i in issues) else "WARNING" if issues else "PASS"


def color_row(ws, row, severity, end=None):
    fill = RED if severity == "FAIL" else AMBER if severity == "WARNING" else GRAY if severity == "EXCLUDED_AR" else None
    if fill:
        for cell in ws[row][:end or ws.max_column]:
            cell.fill = PatternFill("solid", fgColor=fill)


def links(ws, row, col, label, target):
    cell = put(ws, row, col, label)
    cell.hyperlink = f"#'{target[0]}'!A{target[1]}"
    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")


def base_style(ws, widths, freeze):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = False
    ws.sheet_properties.outlinePr.summaryRight = False
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def make_workbook(result: Result) -> Workbook:
    wb = Workbook()
    raw = wb.active
    raw.title = "Raw_INV"
    lookup = wb.create_sheet("Look_UP")
    report = wb.create_sheet("Validation_Report")
    wb.properties.title = "Gary Invoice Lookup"
    wb.properties.subject = f"{result.mode} — {result.status}"
    wb.properties.description = "USD. Snapshot at query time; rerun after changing source data."
    wb.properties.creator = "Gary Invoice Lookup"
    heading(raw, 1, RAW_HEADERS + ["Source Row", "Validation Status", "Validation Message", "Normalized DO"])
    source_to_output = {}
    row_issues = {r.source_row: [i for i in result.issues if r.source_row in i.source_rows] for r in result.source.rows}
    for r in result.source.rows:
        issues = row_issues[r.source_row]
        row_status = status(issues)
        if r.ar and row_status == "PASS":
            row_status = "EXCLUDED_AR"
        values = list(r.values)
        values[3], values[5] = r.selected or None, r.customer or None
        row = append(raw, values + [r.source_row, STATUS_LABELS.get(row_status, row_status), "\n".join(dict.fromkeys(f"[{i.code}] {i.message}" for i in issues)) or ("PRO 以 - AR 结尾，无需 AP 发票。" if r.ar else ""), r.do])
        source_to_output[r.source_row] = row
        raw.cell(row, 2).number_format = '#,##0.00;[Red](#,##0.00);0.00'
        for col in (1, 3, 4, 5, 11):
            raw.cell(row, col).number_format = "@"
        raw.row_dimensions[row].height = min(105, max(30, 16 * (len(issues) + 1)))
        color_row(raw, row, row_status)
    table(raw, 1, raw.max_row, 11, "RawInvoices")
    base_style(raw, [25, 15, 28, 24, 24, 43, 26, 13, 19, 88, 26], "D2")
    raw.print_title_rows = "1:1"

    heading(lookup, 1, LOOKUP_HEADERS + ["Selected", "Selection Reason", "Validation Status", "Validation Message", "Raw Source Rows"])
    for inv in result.invoices:
        issues = [i for i in result.issues if inv.index + 2 in i.lookup_rows]
        reason = ("数据库记录冲突或字段无效" if inv.blocked else "重复数据库记录，不重复计金额" if inv.duplicate else
                  "创建人不在允许名单" if not inv.qualified else
                  ("已按 PRO + IWT 专用客户唯一匹配" if inv.iwt else "已按 DO 客户组/PRO 唯一匹配") if inv.used_by else
                  "未采用：客户组不匹配、候选不唯一或源数据/查询不完整")
        values = list(inv.values)
        values[0], values[3] = inv.pro, inv.number
        if isinstance(values[2], str):
            try:
                values[2] = datetime.fromisoformat(values[2])
            except ValueError:
                pass
        if isinstance(values[2], datetime) and values[2].tzinfo:
            values[2] = values[2].replace(tzinfo=None)
        row = append(lookup, values + ["YES" if inv.used_by else "NO", reason, status(issues), "\n".join(dict.fromkeys(f"[{i.code}] {i.message}" for i in issues)), ", ".join(map(str, sorted(inv.used_by)))])
        lookup.cell(row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
        lookup.cell(row, 5).number_format = '#,##0.00;[Red](#,##0.00);0.00'
        lookup.cell(row, 1).number_format = lookup.cell(row, 4).number_format = "@"
        lookup.row_dimensions[row].height = min(105, max(30, 16 * (len(issues) + 1)))
        color_row(lookup, row, status(issues))
    table(lookup, 1, lookup.max_row, 12, "ScInvoices")
    base_style(lookup, [25, 23, 23, 24, 16, 46, 46, 13, 64, 19, 88, 26], "B2")
    lookup.print_title_rows = "1:1"

    heading(report, 1, ["Gary Invoice Lookup — Validation Report", "Value"])
    pro_set = {r.pro_key for r in result.source.rows if r.pro}
    ar_set = {r.pro_key for r in result.source.rows if r.ar}
    counts = Counter(i.severity for i in result.issues)
    recons = result.reconciliations
    known = [r for r in recons if r.expected is not None]
    complete = [r for r in recons if r.status in {"PASS", "FAIL"}]
    digest = hashlib.sha256(result.source.path.read_bytes()).hexdigest()
    summary = [
        ("Overall status", result.status), ("Data source / 数据来源", result.mode),
        ("Run time", result.started.isoformat(timespec="seconds")),
        ("Source file / sheet", f"{result.source.path.name} / {result.source.sheet}"),
        ("Source SHA256", digest), ("SC filter", f"invoice_date > '{result.since} 00:00:00'; invoice_status <> '3'"),
        ("Currency / comparison", "USD；逐笔按分四舍五入，容差 0.00"),
        ("Allowed creators", "USSDC.ORD / JOHNWU.ORD / SHILONG.ORD"),
        ("Source rows / unique PRO / AR PRO", f"{len(result.source.rows)} / {len(pro_set)} / {len(ar_set)}"),
        ("AP mapped rows / expected AP rows", f"{sum(bool(r.selected) for r in result.source.rows)} / {sum(not r.ar for r in result.source.rows)}"),
        ("Selected distinct invoices", len({key(i.number) for i in result.invoices if i.used_by})),
        ("XC keys completed / incomplete", f"{len(result.xc.succeeded)} / {len(result.xc.failed)}"),
        ("SC keys completed / incomplete", f"{len(result.sc.succeeded)} / {len(result.sc.failed)}"),
        ("FAIL / WARNING exceptions", f"{counts['FAIL']} / {counts['WARNING']}"),
        ("Reconciliation result counts", "; ".join(f"{STATUS_LABELS.get(k,k)}: {v}" for k, v in Counter(r.status for r in recons).items())),
        (f"Known source total ({len(known)} PRO; incomplete sources excluded)", sum((r.expected for r in known), Decimal("0"))),
        (f"Comparable source total ({len(complete)} fully matched PRO)", sum((r.expected for r in complete), Decimal("0"))),
        ("Comparable selected invoice total (same PRO population)", sum((r.actual for r in complete), Decimal("0"))),
        ("Report usage", "本次查询快照。修正源文件或数据库后重新运行；修改导出文件不会重新校验。"),
    ]
    for name, value in summary:
        row = append(report, [name, value])
        report.row_dimensions[row].height = 30 if not isinstance(value, str) or len(value) < 90 else 45
        if isinstance(value, Decimal):
            report.cell(row, 2).number_format = '#,##0.00;[Red](#,##0.00);0.00'
    color_row(report, 2, result.status, 2)
    start = report.max_row + 3
    headers = ["PRO", "Result", "Raw AMT values", "Raw total (USD)", "Selected total (USD)", "Difference (USD)", "Selected invoice#", "Source rows", "Raw link"]
    heading(report, start, headers)
    for item in recons:
        row = append(report, [item.pro, STATUS_LABELS[item.status], item.raw_amounts, item.expected, item.actual, item.difference, item.invoices, ", ".join(map(str, item.source_rows)), ""])
        if item.source_rows:
            links(report, row, 9, "查看 Raw", ("Raw_INV", source_to_output[item.source_rows[0]]))
        for col in (4, 5, 6):
            report.cell(row, col).number_format = '#,##0.00;[Red](#,##0.00);0.00'
        report.row_dimensions[row].height = 36
        color_row(report, row, "FAIL" if item.status in {"INCOMPLETE", "UNDETERMINED", "FAIL"} else item.status, 9)
    table(report, start, report.max_row, len(headers), "ProReconciliation")

    start = report.max_row + 3
    headers = ["Severity", "Check code", "PRO", "DO#", "Source sheet", "Source rows", "Invoice#", "Look_UP rows", "Raw value", "Actual value", "Explanation", "Next action", "Raw link", "Look_UP link"]
    heading(report, start, headers)
    for issue in result.issues:
        row = append(report, [issue.severity, issue.code, issue.pro, issue.dos, result.source.sheet, ", ".join(map(str, issue.source_rows)), issue.invoices, ", ".join(map(str, issue.lookup_rows)), issue.raw_value, issue.actual_value, issue.message, issue.action, "", ""])
        if issue.source_rows:
            links(report, row, 13, "查看 Raw", ("Raw_INV", source_to_output[issue.source_rows[0]]))
        if issue.lookup_rows:
            links(report, row, 14, "查看发票", ("Look_UP", issue.lookup_rows[0]))
        color_row(report, row, issue.severity, 14)
        report.row_dimensions[row].height = 60
    if not result.issues:
        append(report, ["PASS", "NO_EXCEPTIONS", "", "", "", "", "", "", "", "", "没有发现异常。"])
    table(report, start, report.max_row, len(headers), "ValidationExceptions")
    base_style(report, [44, 62, 27, 29, 25, 28, 45, 25, 35, 35, 76, 55, 16, 16], "C2")
    for ws in wb:
        ws.sheet_properties.tabColor = NAVY if ws is not report else "C0504D" if result.status == "FAIL" else "BF9000" if result.status == "WARNING" else NAVY
        ws.sheet_view.zoomScale = 85
    if result.status == "FAIL" or result.mode != "LIVE DATABASE":
        wb.active = 2
    return wb


def export_result(result: Result, output_dir: Path) -> Path:
    output_dir = output_dir.resolve()
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise LookupError("无法创建输出目录，请选择有写入权限的目录。") from exc
    wb = make_workbook(result)
    path = output_dir / f"{result.source.path.stem}_lookup_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    created = False
    try:
        with path.open("xb") as stream:
            created = True
            wb.save(stream)
    except Exception as exc:
        if created:
            path.unlink(missing_ok=True)
        raise LookupError("无法保存结果，请检查输出目录权限和磁盘空间。") from exc
    finally:
        wb.close()
    return path
