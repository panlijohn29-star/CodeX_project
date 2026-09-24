#!/usr/bin/env python3
"""Build DO/IWT delivery allocation workbooks from a carrier invoice list."""

from __future__ import annotations

import io
import os
import re
import sys
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

try:
    import pymysql
except ImportError:  # pragma: no cover - handled at runtime with a clear message.
    pymysql = None  # type: ignore[assignment]

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from platform_config import BASE_DIR, env_int, env_value, get_db_config


MONEY_UNIT = Decimal("0.01")
ZERO = Decimal("0")
DB_CHUNK_SIZE = 500
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
OUTPUT_DIR = Path(BASE_DIR) / "reports_v4" / "ord_do_allocation"

FEATURE = {
    "id": "ord_do_allocation",
    "title": "ORD_DO_ALLOCATION",
    "category": "Tools",
    "description": "Allocate carrier PRO costs across ORD DO/IWT HBL rows and generate AP/AR upload workbooks.",
    "supports_cancel": False,
    "output_type": "interactive",
    "input_schema": [],
    "template": "ord_do_allocation.html",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
INFO_FILL = PatternFill("solid", fgColor="D9EAF7")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN_GRAY = Side(style="thin", color="D9E1F2")
STANDARD_BORDER = Border(bottom=THIN_GRAY)

HIGHLIGHT_CUSTOMERS = {
    "corning incorporated",
    "acer service corporation",
    "on ag",
    "kuehne & nagel inc - ord",
    "kn - aws (rack)",
}

AR_BILLING_PARTIES = {
    "corning incorporated": "APEX LOGISTICS INTERNATIONAL (NY), INC.",
    "acer service corporation": "APEX LOGISTICS INTERNATIONAL (SFO) INC.",
}
HANDLING_OFFICE_AR_PARTIES = {
    "apex-lax": "APEX LOGISTICS INTERNATIONAL (LAX), INC.",
    "apex-jfk": "APEX LOGISTICS INTERNATIONAL (NY), INC.",
    "apex-lck": "APEX LOGISTICS INTERNATIONAL INC.- LCK BRANCH",
    "apex-dfw": "APEX CARGO INTERNATIONAL (DFW) INC.",
    "apex-sfo": "APEX LOGISTICS INTERNATIONAL (SFO) INC.",
}
DO_UPLOAD_BILLING_PARTIES = {
    "old dominion freight c/o k+n ltl ecommerce": "RE TRANS FREIGHT, INC.",
}

RAW_HEADER_ALIASES = {
    "pro": {
        "pro",
        "pro#",
        "pronumber",
        "vendorinvoice",
        "vendorinvoicenumber",
        "invoice",
        "invoiceno",
    },
    "amt": {"amt", "amount", "truckcost", "billingamount", "apamount"},
    "do": {"do", "do#", "dono", "deliveryorder", "dispatchno", "iwt", "iwt#"},
    "customer": {"customer", "customername", "rawcustomer"},
}

OUTPUT_HEADERS = [
    "DO#",
    "PRO#",
    "Trucking company",
    "TRUCK COST",
    "DO COST",
    "customer",
    "delivery location",
    "Handling office",
]
CALC_HEADERS = [
    "DO#",
    "H#",
    "h.cw",
    "DO.cw",
    "DO COST",
    "PRO#",
    "Truck Cost",
    "EST COST",
    "ACTUAL COST",
    "SC_JOB_NO",
    "customer",
    "SHARE WITH OTHER CUSTOMER",
    "Handling office",
]
DO_HEADERS = [
    "DO#",
    "h#",
    "JOB_NO",
    "h.cw",
    "DO.cw",
    "DO COST",
    "EST COST",
    "ACTUAL COST",
    "PRO#",
    "customer",
    "Trucking company",
]
UPLOAD_HEADERS = [
    "AP# (Credit Note#)",
    "OFFICE",
    "Account(billing party)",
    "Vendor Invoice Date",
    "Vendor Invoice Number",
    "JOB NO",
    "RP TYPE",
    "Charge display name",
    "Charge Category ",
    "Charge Code",
    "Amount",
    "UserName",
]
IWT_HEADERS = [
    "DO#",
    "h#",
    "SC_JOB_NO",
    "h.cw",
    "DO.cw",
    "DO COST",
    "EST COST",
    "ACTUAL COST",
    "to",
    "pro#",
    "Trucking company",
    "customer",
]
TMS_HEADERS = [
    "iwt",
    "h#",
    "Pro#",
    "h.cw",
    "DO.cw",
    "DO COST",
    "EST COST",
    "ACTUAL COST",
    "remark",
    "customer",
]
VALIDATION_HEADERS = [
    "Severity",
    "Sheet",
    "Row",
    "PRO#",
    "DO#",
    "H#",
    "Issue Type",
    "Message",
    "Raw Value",
]


class ConverterError(RuntimeError):
    """A fatal, user-actionable conversion error."""


@dataclass
class RawRow:
    source_row: int
    do_no: str
    pro: str
    amount: Optional[Decimal]
    raw_customer: str = ""


@dataclass
class DispatchMeta:
    do_no: str
    trucking_company: str = ""
    customer: str = ""
    delivery_location: str = ""
    handling_office: str = ""


@dataclass
class CarrierLookup:
    by_do: dict[str, str]
    by_pro: dict[str, str]


@dataclass
class CalcRow:
    do_no: str
    hbl_no: str
    h_cw: Optional[Decimal]
    do_cw: Optional[Decimal]
    source_type: str
    raw_row: Optional[RawRow] = None
    pro: str = ""
    truck_cost: Optional[Decimal] = None
    est_cost: Optional[Decimal] = None
    actual_cost: Optional[Decimal] = None
    do_cost: Optional[Decimal] = None
    sc_job_no: str = ""
    customer: str = ""
    share_other_customer: str = ""
    handling_office: str = ""


@dataclass
class ValidationItem:
    severity: str
    sheet: str
    row: str
    pro: str
    do_no: str
    hbl_no: str
    issue_type: str
    message: str
    raw_value: str = ""


class ValidationLog:
    def __init__(self) -> None:
        self.items: list[ValidationItem] = []

    def add(
        self,
        severity: str,
        sheet: str,
        issue_type: str,
        message: str,
        *,
        row: Any = "",
        pro: str = "",
        do_no: str = "",
        hbl_no: str = "",
        raw_value: Any = "",
    ) -> None:
        self.items.append(
            ValidationItem(
                severity=severity,
                sheet=sheet,
                row=clean_text(row),
                pro=pro,
                do_no=do_no,
                hbl_no=hbl_no,
                issue_type=issue_type,
                message=message,
                raw_value=clean_text(raw_value),
            )
        )

    def counts(self) -> Counter[str]:
        return Counter(item.severity for item in self.items)


def configure_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            try:
                reconfigure(encoding="utf-8", errors="backslashreplace")
            except (LookupError, OSError):
                reconfigure(errors="backslashreplace")


def progress(message: str) -> None:
    print(message, flush=True)


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().casefold())


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, Decimal) and value == value.to_integral_value():
        text = str(value.to_integral_value())
    else:
        text = str(value).strip()
    return ILLEGAL_CHARACTERS_RE.sub("", text)


def parse_decimal(value: Any) -> Optional[Decimal]:
    if value is None or isinstance(value, bool):
        return None
    text = clean_text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def money(value: Decimal) -> Decimal:
    return value.quantize(MONEY_UNIT, rounding=ROUND_HALF_UP)


def raw_pro_amount_conflicts(raw_rows: Sequence[RawRow]) -> dict[str, list[RawRow]]:
    grouped: dict[str, list[RawRow]] = defaultdict(list)
    for row in raw_rows:
        grouped[row.pro.casefold()].append(row)
    return {
        pro_key: rows
        for pro_key, rows in grouped.items()
        if len({money(row.amount) for row in rows if row.amount is not None}) > 1
    }


def raw_do_mapping_conflicts(raw_rows: Sequence[RawRow]) -> dict[str, list[RawRow]]:
    grouped: dict[str, list[RawRow]] = defaultdict(list)
    for row in raw_rows:
        grouped[row.do_no.casefold()].append(row)
    return {
        do_key: rows
        for do_key, rows in grouped.items()
        if len(
            {
                (row.pro.casefold(), money(row.amount))
                for row in rows
                if row.amount is not None
            }
        ) > 1
    }


def excel_number(value: Optional[Decimal]) -> Optional[float]:
    if value is None:
        return None
    return float(value)


def get_db_field(row: Any, field_name: str) -> Any:
    if not hasattr(row, "items"):
        return None
    if field_name in row:
        return row[field_name]
    expected = field_name.casefold()
    for key, value in row.items():
        if str(key).casefold() == expected:
            return value
    return None


def chunks(values: Sequence[str], size: int = DB_CHUNK_SIZE) -> Iterable[Sequence[str]]:
    for start in range(0, len(values), size):
        yield values[start : start + size]


def unique_preserve(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value.casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def find_header_row(ws: Worksheet) -> tuple[int, dict[str, int]]:
    normalized_aliases = {
        field: {normalize_header(alias) for alias in aliases}
        for field, aliases in RAW_HEADER_ALIASES.items()
    }
    for row_number in range(1, min(ws.max_row, 25) + 1):
        mapping: dict[str, int] = {}
        for cell in ws[row_number]:
            header = normalize_header(cell.value)
            if not header:
                continue
            for field, aliases in normalized_aliases.items():
                if header in aliases and field not in mapping:
                    mapping[field] = cell.column
        if {"pro", "amt", "do"}.issubset(mapping):
            return row_number, mapping
    raise ConverterError(
        "Cannot find source headers. Required fields: PRO#/vendor invoice number, AMT, DO#."
    )


def read_raw_rows(input_path: Path, validation: ValidationLog) -> list[RawRow]:
    try:
        workbook = load_workbook(input_path, data_only=True, read_only=False)
    except Exception as exc:
        raise ConverterError(f"Cannot open input workbook: {input_path}") from exc

    chosen: Optional[tuple[Worksheet, int, dict[str, int]]] = None
    for ws in workbook.worksheets:
        try:
            chosen = (ws, *find_header_row(ws))
            break
        except ConverterError:
            continue
    if chosen is None:
        raise ConverterError("No worksheet contains the required Raw headers.")

    ws, header_row, mapping = chosen
    rows: list[RawRow] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        do_no = clean_text(ws.cell(row_number, mapping["do"]).value)
        pro = clean_text(ws.cell(row_number, mapping["pro"]).value)
        amount_raw = ws.cell(row_number, mapping["amt"]).value
        amount = parse_decimal(amount_raw)
        customer = (
            clean_text(ws.cell(row_number, mapping["customer"]).value)
            if "customer" in mapping
            else ""
        )
        if not do_no and not pro and amount is None:
            continue
        if not do_no or not pro or amount is None:
            validation.add(
                "ERROR",
                "Raw",
                "RAW_MISSING_REQUIRED",
                "Raw row is missing DO#, PRO#, or a numeric AMT.",
                row=row_number,
                pro=pro,
                do_no=do_no,
                raw_value=amount_raw,
            )
            continue
        rows.append(
            RawRow(
                source_row=row_number,
                do_no=do_no,
                pro=pro,
                amount=amount,
                raw_customer=customer,
            )
        )

    if not rows:
        raise ConverterError("Raw workbook has headers, but no usable data rows.")

    by_do = Counter(row.do_no.casefold() for row in rows)
    for row in rows:
        if by_do[row.do_no.casefold()] > 1:
            validation.add(
                "WARNING",
                "Raw",
                "RAW_DUPLICATE_DO",
                "The same DO# appears more than once in Raw. The first mapping is used for lookup fields.",
                row=row.source_row,
                pro=row.pro,
                do_no=row.do_no,
            )

    for conflict_rows in raw_pro_amount_conflicts(rows).values():
        first = conflict_rows[0]
        details = "; ".join(
            f"row {row.source_row}={money(row.amount):.2f}"
            for row in conflict_rows
            if row.amount is not None
        )
        validation.add(
            "ERROR",
            "Raw",
            "RAW_PRO_AMOUNT_CONFLICT",
            "The same PRO# has different Raw AMT values. Allocation uses the first amount, so the source data requires review.",
            row=first.source_row,
            pro=first.pro,
            do_no=first.do_no,
            raw_value=details,
        )

    for conflict_rows in raw_do_mapping_conflicts(rows).values():
        first = conflict_rows[0]
        details = "; ".join(
            f"row {row.source_row}: PRO# {row.pro}, AMT {money(row.amount):.2f}"
            for row in conflict_rows
            if row.amount is not None
        )
        validation.add(
            "ERROR",
            "Raw",
            "RAW_DO_MAPPING_CONFLICT",
            "The same DO# has different PRO#/AMT mappings. Lookup and allocation use the first mapping, so the source data requires review.",
            row=first.source_row,
            pro=first.pro,
            do_no=first.do_no,
            raw_value=details,
        )
    return rows


def read_carrier_lookup(carrier_path: Optional[Path], validation: ValidationLog) -> CarrierLookup:
    if carrier_path is None or not carrier_path.exists():
        return CarrierLookup(by_do={}, by_pro={})

    try:
        workbook = load_workbook(carrier_path, data_only=True, read_only=False)
    except Exception as exc:
        validation.add(
            "WARNING",
            "Carrier",
            "CARRIER_FILE_UNREADABLE",
            f"Carrier lookup file could not be opened: {carrier_path}",
            raw_value=exc,
        )
        return CarrierLookup(by_do={}, by_pro={})

    by_do: dict[str, str] = {}
    by_pro: dict[str, str] = {}
    for ws in workbook.worksheets:
        header_row = None
        mapping: dict[str, int] = {}
        for row_number in range(1, min(ws.max_row, 10) + 1):
            current: dict[str, int] = {}
            for cell in ws[row_number]:
                header = normalize_header(cell.value)
                if header in {"iwt", "iwtno", "iwt#"}:
                    current["do"] = cell.column
                elif header in {"pro", "pro#", "prono"}:
                    current["pro"] = cell.column
                elif header in {"carrier", "truckingcompany", "truckcompany", "vendor"}:
                    current["carrier"] = cell.column
            if {"pro", "carrier"}.issubset(current):
                header_row = row_number
                mapping = current
                break
        if header_row is None:
            continue
        for row_number in range(header_row + 1, ws.max_row + 1):
            carrier = clean_text(ws.cell(row_number, mapping["carrier"]).value)
            if not carrier:
                continue
            pro = clean_text(ws.cell(row_number, mapping["pro"]).value)
            do_no = clean_text(ws.cell(row_number, mapping["do"]).value) if "do" in mapping else ""
            if do_no:
                by_do.setdefault(do_no.casefold(), carrier)
            if pro:
                by_pro.setdefault(pro.casefold(), carrier)

    if not by_do and not by_pro:
        validation.add(
            "WARNING",
            "Carrier",
            "CARRIER_FILE_NO_MATCH",
            f"Carrier lookup file has no recognizable pro#/carrier rows: {carrier_path}",
            raw_value=carrier_path,
        )
    return CarrierLookup(by_do=by_do, by_pro=by_pro)


def env_bool(name: str, default: bool) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().casefold() not in {"0", "false", "no", "off"}


def load_config(env_path: Optional[Path] = None, script_dir: Optional[Path] = None) -> tuple[dict[str, Any], dict[str, Any]]:
    if pymysql is None:
        raise ConverterError("PyMySQL is not installed. Run pip install -r requirements.txt.")
    xc_values = {
        "host": env_value("XC_DB_HOST", env_value("XC_HOST")),
        "user": env_value("XC_DB_USER", env_value("XC_USER")),
        "password": env_value("XC_DB_PASSWORD", env_value("XC_PASSWORD")),
        "db": env_value("XC_DB_NAME", env_value("XC_DATABASE", env_value("XC_DB"))),
    }
    missing = [name for name, value in xc_values.items() if not value]
    if missing:
        raise ConverterError("XC database configuration is missing: {0}. Add XC_DB_HOST, XC_DB_PORT, XC_DB_USER, XC_DB_PASSWORD, and XC_DB_NAME to the Platform .env file.".format(", ".join(missing)))
    xc_config = {
        **xc_values, "port": env_int("XC_DB_PORT", env_int("XC_PORT", 3306)), "charset": env_value("XC_DB_CHARSET", "utf8mb4"),
        "connect_timeout": env_int("DB_CONNECT_TIMEOUT", 10), "read_timeout": env_int("DB_READ_TIMEOUT", 600), "write_timeout": env_int("DB_WRITE_TIMEOUT", 600),
    }
    try:
        sc_config = get_db_config("scdbus")
    except RuntimeError as exc:
        raise ConverterError("SCDBUS database configuration is incomplete: {0}".format(exc)) from exc
    for config in (xc_config, sc_config):
        config["cursorclass"] = pymysql.cursors.DictCursor
        config["autocommit"] = True
    return xc_config, sc_config


def fetch_all(connection_config: dict[str, Any], sql: str, params: Sequence[Any]) -> list[dict[str, Any]]:
    if pymysql is None:
        raise ConverterError("PyMySQL is not installed. Run pip install -r requirements.txt.")
    connection = pymysql.connect(**connection_config)
    try:
        with connection.cursor() as cursor:
            cursor.execute(sql, tuple(params))
            return list(cursor.fetchall())
    finally:
        connection.close()


def query_dispatch_meta(xc_config: dict[str, Any], do_numbers: Sequence[str]) -> dict[str, DispatchMeta]:
    result: dict[str, DispatchMeta] = {}
    for group in chunks(do_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                driver_company_code AS trucking_company,
                customer_code AS customer,
                delivery_location,
                dest_handing_office AS handling_office,
                dispatch_no
            FROM t_dispatch_head
            WHERE dispatch_no IN ({placeholders})
        """
        for row in fetch_all(xc_config, sql, group):
            do_no = clean_text(get_db_field(row, "dispatch_no"))
            if not do_no:
                continue
            result[do_no.casefold()] = DispatchMeta(
                do_no=do_no,
                trucking_company=clean_text(get_db_field(row, "trucking_company")),
                customer=clean_text(get_db_field(row, "customer")),
                delivery_location=clean_text(get_db_field(row, "delivery_location")),
                handling_office=clean_text(get_db_field(row, "handling_office")),
            )
    return result


def query_iwt_meta(xc_config: dict[str, Any], iwt_numbers: Sequence[str]) -> dict[str, DispatchMeta]:
    result: dict[str, DispatchMeta] = {}
    for group in chunks(iwt_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                iwt_no,
                driver_company_code AS trucking_company
            FROM t_iwt_head
            WHERE iwt_no IN ({placeholders})
        """
        for row in fetch_all(xc_config, sql, group):
            iwt_no = clean_text(get_db_field(row, "iwt_no"))
            if not iwt_no:
                continue
            result[iwt_no.casefold()] = DispatchMeta(
                do_no=iwt_no,
                trucking_company=clean_text(get_db_field(row, "trucking_company")),
            )
    return result


def query_xc_dispatch_calc(xc_config: dict[str, Any], do_numbers: Sequence[str]) -> list[CalcRow]:
    rows: list[CalcRow] = []
    for group in chunks(do_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                head.dispatch_no,
                body.house_no,
                total.chargeable_weight,
                sum_table.do_chargeable_weight_sum
            FROM t_dispatch_head head
            LEFT JOIN t_dispatch_body body
                ON body.dispatch_head_id = head.id
            LEFT JOIN t_shipment_total total
                ON total.bussiness_no = body.house_no
               AND total.data_type = 'H'
            LEFT JOIN (
                SELECT
                    h.dispatch_no,
                    SUM(t.chargeable_weight) AS do_chargeable_weight_sum
                FROM t_dispatch_head h
                LEFT JOIN t_dispatch_body b
                    ON b.dispatch_head_id = h.id
                LEFT JOIN t_shipment_total t
                    ON t.bussiness_no = b.house_no
                   AND t.data_type = 'H'
                WHERE h.dispatch_no IN ({placeholders})
                GROUP BY h.dispatch_no
            ) sum_table
                ON sum_table.dispatch_no = head.dispatch_no
            WHERE head.dispatch_no IN ({placeholders})
        """
        params = list(group) + list(group)
        for row in fetch_all(xc_config, sql, params):
            rows.append(
                CalcRow(
                    do_no=clean_text(get_db_field(row, "dispatch_no")),
                    hbl_no=clean_text(get_db_field(row, "house_no")),
                    h_cw=parse_decimal(get_db_field(row, "chargeable_weight")),
                    do_cw=parse_decimal(get_db_field(row, "do_chargeable_weight_sum")),
                    source_type="DO",
                )
            )
    return rows


def query_xc_iwt_calc(xc_config: dict[str, Any], do_numbers: Sequence[str]) -> list[CalcRow]:
    rows: list[CalcRow] = []
    for group in chunks(do_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                head.iwt_no,
                body.house_no,
                total.chargeable_weight,
                sum_table.do_chargeable_weight_sum
            FROM t_iwt_head head
            LEFT JOIN t_iwt_body body
                ON body.iwt_head_id = head.id
            LEFT JOIN t_shipment_total total
                ON total.bussiness_no = body.house_no
               AND total.data_type = 'H'
            LEFT JOIN (
                SELECT
                    h.iwt_no,
                    SUM(t.chargeable_weight) AS do_chargeable_weight_sum
                FROM t_iwt_head h
                LEFT JOIN t_iwt_body b
                    ON b.iwt_head_id = h.id
                LEFT JOIN t_shipment_total t
                    ON t.bussiness_no = b.house_no
                   AND t.data_type = 'H'
                WHERE h.iwt_no IN ({placeholders})
                GROUP BY h.iwt_no
            ) sum_table
                ON sum_table.iwt_no = head.iwt_no
            WHERE head.iwt_no IN ({placeholders})
        """
        params = list(group) + list(group)
        for row in fetch_all(xc_config, sql, params):
            rows.append(
                CalcRow(
                    do_no=clean_text(get_db_field(row, "iwt_no")),
                    hbl_no=clean_text(get_db_field(row, "house_no")),
                    h_cw=parse_decimal(get_db_field(row, "chargeable_weight")),
                    do_cw=parse_decimal(get_db_field(row, "do_chargeable_weight_sum")),
                    source_type="IWT",
                )
            )
    return rows


def query_sc_fallback(sc_config: dict[str, Any], do_numbers: Sequence[str]) -> list[CalcRow]:
    rows: list[CalcRow] = []
    for group in chunks(do_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                MAWB_NO,
                HAWB_NO,
                MAWB_CHARGEABLE_WEIGHT AS h_CW,
                (
                    SELECT SUM(MAWB_CHARGEABLE_WEIGHT)
                    FROM op_ae_job ABC
                    LEFT JOIN op_ae_job_booking AB
                        ON ABC.JOB_ID = AB.JOB_ID
                    WHERE AB.MAWB_NO = op_ae_job_booking.MAWB_NO
                      AND ABC.JOB_MODE <> '02'
                ) AS DO_CW
            FROM op_ae_job_booking
            LEFT JOIN OP_AE_JOB
                ON op_ae_job_booking.JOB_ID = OP_AE_JOB.JOB_ID
            WHERE MAWB_NO IN ({placeholders})
              AND JOB_MODE <> '02'
        """
        for row in fetch_all(sc_config, sql, group):
            rows.append(
                CalcRow(
                    do_no=clean_text(get_db_field(row, "MAWB_NO")),
                    hbl_no=clean_text(get_db_field(row, "HAWB_NO")),
                    h_cw=parse_decimal(get_db_field(row, "h_CW")),
                    do_cw=parse_decimal(get_db_field(row, "DO_CW")),
                    source_type="SC",
                )
            )
    return rows


def query_sc_jobinfo_by_company_field(
    sc_config: dict[str, Any],
    hbl_numbers: Sequence[str],
    company_field: str,
) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    for group in chunks(hbl_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT
                hbl_no,
                job_no AS SC_JOB_NO,
                customer
            FROM v_jobinfo
            WHERE hbl_no IN ({placeholders})
              AND {company_field} = 'apex-ord'
        """
        for row in fetch_all(sc_config, sql, group):
            hbl_no = clean_text(get_db_field(row, "hbl_no"))
            if not hbl_no:
                continue
            result.setdefault(
                hbl_no.casefold(),
                (
                    clean_text(get_db_field(row, "SC_JOB_NO")),
                    clean_text(get_db_field(row, "customer")),
                ),
            )
    return result


def query_sc_jobinfo(sc_config: dict[str, Any], hbl_numbers: Sequence[str]) -> dict[str, tuple[str, str]]:
    result = query_sc_jobinfo_by_company_field(sc_config, hbl_numbers, "op_company")
    missing_hbls = [
        hbl_no
        for hbl_no in hbl_numbers
        if hbl_no and hbl_no.casefold() not in result
    ]
    if missing_hbls:
        sales_matches = query_sc_jobinfo_by_company_field(
            sc_config,
            missing_hbls,
            "sales_company",
        )
        for hbl_key, values in sales_matches.items():
            result.setdefault(hbl_key, values)
    return result


def query_iwt_to(xc_config: dict[str, Any], iwt_numbers: Sequence[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for group in chunks(iwt_numbers):
        placeholders = ", ".join(["%s"] * len(group))
        sql = f"""
            SELECT iwt_no AS `DO#`, to_firm_code AS `to`
            FROM t_iwt_head
            WHERE iwt_no IN ({placeholders})
        """
        for row in fetch_all(xc_config, sql, group):
            iwt_no = clean_text(get_db_field(row, "DO#"))
            if iwt_no:
                result[iwt_no.casefold()] = clean_text(get_db_field(row, "to"))
    return result


def apply_raw_mappings(calc_rows: list[CalcRow], raw_rows: Sequence[RawRow]) -> None:
    raw_by_do: dict[str, RawRow] = {}
    for raw in raw_rows:
        raw_by_do.setdefault(raw.do_no.casefold(), raw)
    for row in calc_rows:
        raw = raw_by_do.get(row.do_no.casefold())
        row.raw_row = raw
        if raw:
            row.pro = raw.pro
            row.truck_cost = raw.amount


def apply_handling_office(
    calc_rows: list[CalcRow],
    dispatch_meta: dict[str, DispatchMeta],
) -> None:
    for row in calc_rows:
        meta = dispatch_meta.get(row.do_no.casefold())
        row.handling_office = meta.handling_office if meta else ""


def enrich_from_sc_jobinfo(
    calc_rows: list[CalcRow],
    sc_config: dict[str, Any],
    validation: ValidationLog,
) -> None:
    hbls = unique_preserve(row.hbl_no for row in calc_rows if row.hbl_no)
    jobinfo = query_sc_jobinfo(sc_config, hbls) if hbls else {}
    for row in calc_rows:
        if not row.hbl_no:
            validation.add(
                "ERROR",
                "calculation",
                "MISSING_HBL",
                "Database returned a blank H# for this DO#.",
                pro=row.pro,
                do_no=row.do_no,
            )
            continue
        match = jobinfo.get(row.hbl_no.casefold())
        if not match:
            validation.add(
                "WARNING",
                "calculation",
                "SC_JOB_NOT_FOUND",
                "SC v_jobinfo did not return a job/customer for this H#.",
                pro=row.pro,
                do_no=row.do_no,
                hbl_no=row.hbl_no,
                raw_value=row.hbl_no,
            )
            continue
        row.sc_job_no, row.customer = match


def allocate_costs(calc_rows: list[CalcRow], validation: ValidationLog) -> None:
    by_pro: dict[str, list[CalcRow]] = defaultdict(list)
    for row in calc_rows:
        if row.pro:
            by_pro[row.pro.casefold()].append(row)

    for pro_key, rows in by_pro.items():
        pro = rows[0].pro
        amount = next((row.truck_cost for row in rows if row.truck_cost is not None), None)
        if amount is None:
            validation.add(
                "ERROR",
                "calculation",
                "MISSING_PRO_AMOUNT",
                "No numeric Raw AMT is available for this PRO#.",
                pro=pro,
            )
            continue
        valid_rows = [row for row in rows if row.h_cw is not None and row.h_cw > ZERO]
        total_cw = sum((row.h_cw for row in valid_rows if row.h_cw is not None), ZERO)
        if total_cw <= ZERO:
            validation.add(
                "ERROR",
                "calculation",
                "INVALID_PRO_CW",
                "No positive chargeable weight exists for this PRO#, so cost cannot be allocated.",
                pro=pro,
            )
            for row in rows:
                row.est_cost = ZERO
                row.actual_cost = ZERO
            continue

        for row in rows:
            if row.h_cw is None or row.h_cw <= ZERO:
                row.est_cost = ZERO
                row.actual_cost = ZERO
                validation.add(
                    "ERROR",
                    "calculation",
                    "INVALID_H_CW",
                    "H# has blank, zero, or negative chargeable weight. It received 0 allocation.",
                    pro=row.pro,
                    do_no=row.do_no,
                    hbl_no=row.hbl_no,
                    raw_value=row.h_cw,
                )
            else:
                row.est_cost = money(row.h_cw / total_cw * amount)
                row.actual_cost = row.est_cost

        actual_sum = sum((row.actual_cost or ZERO for row in rows), ZERO)
        diff = actual_sum - amount
        if diff != ZERO and valid_rows:
            first = valid_rows[0]
            first.actual_cost = money((first.actual_cost or ZERO) - diff)

    by_do: dict[str, list[CalcRow]] = defaultdict(list)
    for row in calc_rows:
        by_do[row.do_no.casefold()].append(row)
    for rows in by_do.values():
        do_total = sum((row.actual_cost or ZERO for row in rows), ZERO)
        for row in rows:
            row.do_cost = money(do_total)

    customers_by_pro: dict[str, set[str]] = defaultdict(set)
    for row in calc_rows:
        if row.pro and row.customer:
            customers_by_pro[row.pro.casefold()].add(row.customer.casefold())
    for row in calc_rows:
        row.share_other_customer = "Y" if len(customers_by_pro.get(row.pro.casefold(), set())) > 1 else "N"


def compare_pro_totals(raw_rows: Sequence[RawRow], calc_rows: Sequence[CalcRow], validation: ValidationLog) -> None:
    raw_amount: dict[str, Decimal] = {}
    for raw in raw_rows:
        if raw.amount is not None:
            raw_amount.setdefault(raw.pro.casefold(), raw.amount)

    actual_by_pro: dict[str, Decimal] = defaultdict(lambda: ZERO)
    for row in calc_rows:
        if row.pro:
            actual_by_pro[row.pro.casefold()] += row.actual_cost or ZERO

    for raw in raw_rows:
        expected = raw_amount.get(raw.pro.casefold())
        actual = actual_by_pro.get(raw.pro.casefold(), ZERO)
        if expected is not None and money(actual) != money(expected):
            validation.add(
                "ERROR",
                "output",
                "PRO_TOTAL_MISMATCH",
                f"Sum of allocated DO COST for PRO# {raw.pro} is {actual:.2f}, but Raw AMT is {expected:.2f}.",
                row=raw.source_row,
                pro=raw.pro,
                do_no=raw.do_no,
                raw_value=expected,
            )


def get_trucking_for_pro(
    pro: str,
    raw_rows: Sequence[RawRow],
    dispatch_meta: dict[str, DispatchMeta],
    carrier_lookup: CarrierLookup,
) -> str:
    for raw in raw_rows:
        if raw.pro.casefold() == pro.casefold():
            meta = dispatch_meta.get(raw.do_no.casefold())
            if meta and meta.trucking_company:
                return meta.trucking_company
    carrier = carrier_lookup.by_pro.get(pro.casefold())
    if carrier:
        return carrier
    return ""


def exclude_from_do_sheet(row: CalcRow) -> bool:
    customer = row.customer.casefold()
    if customer == "on ag":
        return True
    return customer == "vertiv corporation" and row.share_other_customer == "N"


def handling_office_ar_party(handling_office: str) -> str:
    return HANDLING_OFFICE_AR_PARTIES.get(handling_office.strip().casefold(), "")


def do_upload_billing_party(trucking_company: str) -> str:
    normalized_company = trucking_company.strip().casefold()
    return DO_UPLOAD_BILLING_PARTIES.get(normalized_company, trucking_company)


def safe_sheet_name(name: str) -> str:
    return name[:31]


def append_table(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))


def style_sheet(ws: Worksheet, money_cols: Sequence[int] = (), numeric_cols: Sequence[int] = ()) -> None:
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = STANDARD_BORDER
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = STANDARD_BORDER
            cell.alignment = Alignment(vertical="top")

    for column in money_cols:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2, max_row=ws.max_row):
            for item in cell:
                item.number_format = '#,##0.00'
    for column in numeric_cols:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2, max_row=ws.max_row):
            for item in cell:
                item.number_format = '#,##0.00'

    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = 10
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            width = max(width, min(len(str(value)) + 2, 42))
        ws.column_dimensions[letter].width = width


def fill_row(ws: Worksheet, row_number: int, fill: PatternFill) -> None:
    for cell in ws[row_number]:
        cell.fill = fill


def write_workbook(
    output_path: Path,
    raw_rows: Sequence[RawRow],
    calc_rows: Sequence[CalcRow],
    dispatch_meta: dict[str, DispatchMeta],
    carrier_lookup: CarrierLookup,
    iwt_to: dict[str, str],
    validation: ValidationLog,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    ws_output = wb.create_sheet("output")
    ws_calc = wb.create_sheet("calculation")
    ws_do = wb.create_sheet("DO")
    ws_iwt = wb.create_sheet("IWT")
    ws_tms = wb.create_sheet("TMS-billing")
    ws_validation = wb.create_sheet("Validation_report")
    ws_do_upload = wb.create_sheet("DO_upload")
    ws_do_arupload = wb.create_sheet("DO_ARupload")

    calc_by_do: dict[str, list[CalcRow]] = defaultdict(list)
    for row in calc_rows:
        calc_by_do[row.do_no.casefold()].append(row)

    output_rows = []
    for raw in raw_rows:
        meta = dispatch_meta.get(raw.do_no.casefold(), DispatchMeta(do_no=raw.do_no))
        trucking_company = (
            meta.trucking_company
            or get_trucking_for_pro(raw.pro, raw_rows, dispatch_meta, carrier_lookup)
            or carrier_lookup.by_do.get(raw.do_no.casefold(), "")
            or carrier_lookup.by_pro.get(raw.pro.casefold(), "")
        )
        do_cost = sum((row.actual_cost or ZERO for row in calc_by_do.get(raw.do_no.casefold(), [])), ZERO)
        output_rows.append(
            [
                raw.do_no,
                raw.pro,
                trucking_company,
                excel_number(raw.amount),
                excel_number(money(do_cost)),
                meta.customer,
                meta.delivery_location,
                meta.handling_office,
            ]
        )
    append_table(ws_output, OUTPUT_HEADERS, output_rows)

    conflicting_pro_keys = set(raw_pro_amount_conflicts(raw_rows))
    conflicting_do_keys = set(raw_do_mapping_conflicts(raw_rows))
    for idx, raw in enumerate(raw_rows, start=2):
        rows_for_pro = [
            row
            for row in calc_rows
            if row.pro.casefold() == raw.pro.casefold()
        ]
        actual = sum((row.actual_cost or ZERO for row in rows_for_pro), ZERO)
        if (
            raw.pro.casefold() in conflicting_pro_keys
            or raw.do_no.casefold() in conflicting_do_keys
            or (raw.amount is not None and money(actual) != money(raw.amount))
        ):
            fill_row(ws_output, idx, ERROR_FILL)
        if raw.do_no.casefold() not in calc_by_do:
            fill_row(ws_output, idx, WARNING_FILL)
        meta = dispatch_meta.get(raw.do_no.casefold())
        if raw.do_no.upper().startswith("DO") and (meta is None or not meta.handling_office):
            validation.add(
                "WARNING",
                "output",
                "HANDLING_OFFICE_MISSING",
                "DO# starts with DO but XC t_dispatch_head.dest_handing_office is blank.",
                row=idx,
                pro=raw.pro,
                do_no=raw.do_no,
            )

    append_table(
        ws_calc,
        CALC_HEADERS,
        [
            [
                row.do_no,
                row.hbl_no,
                excel_number(row.h_cw),
                excel_number(row.do_cw),
                excel_number(row.do_cost),
                row.pro,
                excel_number(row.truck_cost),
                None,
                None,
                row.sc_job_no,
                row.customer,
                row.share_other_customer,
                row.handling_office,
            ]
            for row in calc_rows
        ],
    )
    for excel_row, row in enumerate(calc_rows, start=2):
        ws_calc.cell(excel_row, 8).value = f'=IFERROR(ROUND(C{excel_row}/SUMIF(F:F,F{excel_row},C:C)*G{excel_row},2),0)'
        ws_calc.cell(excel_row, 9).value = (
            f'=IF(SUMIF(F:F,F{excel_row},H:H)=G{excel_row},H{excel_row},'
            f'IF(COUNTIF($F$2:F{excel_row},F{excel_row})=1,'
            f'H{excel_row}-(SUMIF(F:F,F{excel_row},H:H)-G{excel_row}),H{excel_row}))'
        )
        if row.customer.casefold() in HIGHLIGHT_CUSTOMERS:
            fill_row(ws_calc, excel_row, ORANGE_FILL)
        if row.handling_office and row.handling_office.casefold() != "apex-ord":
            fill_row(ws_calc, excel_row, ORANGE_FILL)
        if row.h_cw is None or row.h_cw <= ZERO or not row.hbl_no:
            fill_row(ws_calc, excel_row, ERROR_FILL)

    do_rows = []
    iwt_rows = []
    tms_rows = []
    do_upload_rows = []
    do_ar_rows = []
    for row in calc_rows:
        trucking = get_trucking_for_pro(row.pro, raw_rows, dispatch_meta, carrier_lookup)
        if row.do_no.upper().startswith("DO"):
            handling_party = ""
            if row.handling_office and row.handling_office.casefold() != "apex-ord":
                handling_party = handling_office_ar_party(row.handling_office)
                if not handling_party:
                    validation.add(
                        "WARNING",
                        "DO_ARupload",
                        "HANDLING_OFFICE_AR_MAPPING_MISSING",
                        "Handling office is not APEX-ORD but has no AR billing-party mapping.",
                        pro=row.pro,
                        do_no=row.do_no,
                        hbl_no=row.hbl_no,
                        raw_value=row.handling_office,
                    )
            ar_party = handling_party or AR_BILLING_PARTIES.get(row.customer.casefold(), "")
            if ar_party:
                do_ar_rows.append(
                    [
                        "",
                        "APEX-ORD",
                        ar_party,
                        "",
                        row.do_no,
                        row.sc_job_no,
                        "AR",
                        "Delivery",
                        "Trucking",
                        "8000",
                        excel_number(row.actual_cost),
                        "JOHNWU.ORD",
                    ]
                )
            if exclude_from_do_sheet(row):
                continue
            do_rows.append(
                [
                    row.do_no,
                    row.hbl_no,
                    row.sc_job_no,
                    excel_number(row.h_cw),
                    excel_number(row.do_cw),
                    excel_number(row.do_cost),
                    excel_number(row.est_cost),
                    excel_number(row.actual_cost),
                    row.pro,
                    row.customer,
                    trucking,
                ]
            )
            if row.customer.casefold() != "kuehne & nagel inc - ord":
                do_upload_rows.append(
                    [
                        "",
                        "APEX-ORD",
                        do_upload_billing_party(trucking),
                        "",
                        row.pro,
                        row.sc_job_no,
                        "AP",
                        "Delivery",
                        "Trucking",
                        "8000",
                        excel_number(row.actual_cost),
                        "JOHNWU.ORD",
                    ]
                )
            if row.customer.casefold() == "kuehne & nagel inc - ord":
                tms_rows.append(
                    [
                        row.do_no,
                        row.hbl_no,
                        row.pro,
                        excel_number(row.h_cw),
                        excel_number(row.do_cw),
                        excel_number(row.do_cost),
                        excel_number(row.est_cost),
                        excel_number(row.actual_cost),
                        "",
                        row.customer,
                    ]
                )
        elif row.do_no.upper().startswith("IWT"):
            iwt_rows.append(
                [
                    row.do_no,
                    row.hbl_no,
                    row.sc_job_no,
                    excel_number(row.h_cw),
                    excel_number(row.do_cw),
                    excel_number(row.do_cost),
                    excel_number(row.est_cost),
                    excel_number(row.actual_cost),
                    iwt_to.get(row.do_no.casefold(), ""),
                    row.pro,
                    trucking,
                    row.customer,
                ]
            )

    append_table(ws_do, DO_HEADERS, do_rows)
    append_table(ws_iwt, IWT_HEADERS, iwt_rows)
    append_table(ws_tms, TMS_HEADERS, tms_rows)
    append_table(ws_do_upload, UPLOAD_HEADERS, do_upload_rows)
    append_table(ws_do_arupload, UPLOAD_HEADERS, do_ar_rows)
    calc_last_row = max(ws_calc.max_row, 2)
    for excel_row in range(2, ws_do_arupload.max_row + 1):
        ws_do_arupload.cell(excel_row, 11).value = (
            f'=IFERROR(INDEX(\'calculation\'!$I$2:$I${calc_last_row},'
            f'MATCH(F{excel_row},\'calculation\'!$J$2:$J${calc_last_row},0)),0)'
        )

    append_table(
        ws_validation,
        VALIDATION_HEADERS,
        [
            [
                item.severity,
                item.sheet,
                item.row,
                item.pro,
                item.do_no,
                item.hbl_no,
                item.issue_type,
                item.message,
                item.raw_value,
            ]
            for item in validation.items
        ],
    )
    for excel_row, item in enumerate(validation.items, start=2):
        if item.severity == "ERROR":
            fill_row(ws_validation, excel_row, ERROR_FILL)
        elif item.severity == "WARNING":
            fill_row(ws_validation, excel_row, WARNING_FILL)
        else:
            fill_row(ws_validation, excel_row, INFO_FILL)

    style_sheet(ws_output, money_cols=[4, 5])
    style_sheet(ws_calc, money_cols=[5, 7, 8, 9], numeric_cols=[3, 4])
    style_sheet(ws_do, money_cols=[6, 7, 8], numeric_cols=[4, 5])
    style_sheet(ws_iwt, money_cols=[6, 7, 8], numeric_cols=[4, 5])
    style_sheet(ws_tms, money_cols=[6, 7, 8], numeric_cols=[4, 5])
    style_sheet(ws_do_upload, money_cols=[11])
    style_sheet(ws_do_arupload, money_cols=[11])
    style_sheet(ws_validation)
    ws_validation.column_dimensions["H"].width = 58
    ws_validation.column_dimensions["I"].width = 44
    for row_number in range(2, ws_validation.max_row + 1):
        ws_validation.cell(row_number, 8).alignment = Alignment(vertical="top", wrap_text=True)
        ws_validation.cell(row_number, 9).alignment = Alignment(vertical="top", wrap_text=True)
        ws_validation.row_dimensions[row_number].height = 32

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def choose_output_path(input_path: Path, output_dir: Optional[Path]) -> Path:
    folder = output_dir or input_path.parent
    base = folder / f"{input_path.stem}_allocated.xlsx"
    if not base.exists():
        return base
    from datetime import datetime

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return folder / f"{input_path.stem}_allocated_{stamp}.xlsx"


def build_workbook(
    input_path: Path,
    output_path: Path,
    env_path: Optional[Path],
    carrier_path: Optional[Path] = None,
    use_existing_output: Optional[Path] = None,
) -> tuple[Path, Counter[str]]:
    validation = ValidationLog()
    progress(f"[1/8] Reading Raw workbook: {input_path}")
    raw_rows = read_raw_rows(input_path, validation)
    progress(f"      Found {len(raw_rows)} Raw row(s).")
    carrier_lookup = read_carrier_lookup(carrier_path, validation)

    if use_existing_output:
        progress(f"[2/8] Reading debug/reference workbook: {use_existing_output}")
        # Debug path for layout testing. It uses an existing reference workbook's calculation data.
        ref = load_workbook(use_existing_output, data_only=True)
        calc_rows = []
        if "calculation" not in ref.sheetnames:
            raise ConverterError("--use-existing-output must contain a calculation sheet.")
        ws = ref["calculation"]
        for row_number in range(2, ws.max_row + 1):
            do_no = clean_text(ws.cell(row_number, 1).value)
            hbl = clean_text(ws.cell(row_number, 2).value)
            if not do_no and not hbl:
                continue
            calc_rows.append(
                CalcRow(
                    do_no=do_no,
                    hbl_no=hbl,
                    h_cw=parse_decimal(ws.cell(row_number, 3).value),
                    do_cw=parse_decimal(ws.cell(row_number, 4).value),
                    source_type="REF",
                    sc_job_no=clean_text(ws.cell(row_number, 10).value),
                    customer=clean_text(ws.cell(row_number, 11).value),
                    share_other_customer=clean_text(ws.cell(row_number, 12).value),
                    handling_office=clean_text(ws.cell(row_number, 13).value),
                )
            )
        dispatch_meta = {}
        if "output" in ref.sheetnames:
            ws_out = ref["output"]
            for row_number in range(2, ws_out.max_row + 1):
                do_no = clean_text(ws_out.cell(row_number, 1).value)
                if do_no:
                    dispatch_meta[do_no.casefold()] = DispatchMeta(
                        do_no=do_no,
                        trucking_company=clean_text(ws_out.cell(row_number, 3).value),
                        customer=clean_text(ws_out.cell(row_number, 6).value),
                        delivery_location=clean_text(ws_out.cell(row_number, 7).value),
                        handling_office=clean_text(ws_out.cell(row_number, 8).value),
                    )
        iwt_to = {}
    else:
        progress("[2/8] Loading database configuration...")
        xc_config, sc_config = load_config(env_path, Path(__file__).resolve().parent)
        do_numbers = unique_preserve(row.do_no for row in raw_rows)
        iwt_numbers = [do_no for do_no in do_numbers if do_no.upper().startswith("IWT")]
        dispatch_numbers = [do_no for do_no in do_numbers if not do_no.upper().startswith("IWT")]
        progress(
            f"[3/8] Querying XC dispatch/IWT data "
            f"({len(dispatch_numbers)} dispatch-like DO, {len(iwt_numbers)} IWT)..."
        )
        try:
            dispatch_meta = query_dispatch_meta(xc_config, dispatch_numbers)
            dispatch_meta.update(query_iwt_meta(xc_config, iwt_numbers))
            dispatch_rows = query_xc_dispatch_calc(xc_config, dispatch_numbers)
            iwt_rows = query_xc_iwt_calc(xc_config, iwt_numbers)
        except pymysql.MySQLError as exc:
            raise ConverterError(f"XC database query failed: {exc}") from exc

        found_keys = {
            row.do_no.casefold()
            for row in dispatch_rows + iwt_rows
            if row.do_no
        }
        missing_for_xc = [do_no for do_no in do_numbers if do_no.casefold() not in found_keys]
        progress(
            f"[4/8] Querying SC fallback for {len(missing_for_xc)} DO(s) not found in XC..."
        )
        try:
            sc_rows = query_sc_fallback(sc_config, missing_for_xc)
        except pymysql.MySQLError as exc:
            raise ConverterError(f"SC fallback query failed: {exc}") from exc

        found_after_sc = {row.do_no.casefold() for row in sc_rows if row.do_no}
        for do_no in missing_for_xc:
            if do_no.casefold() not in found_after_sc:
                raw = next((item for item in raw_rows if item.do_no.casefold() == do_no.casefold()), None)
                validation.add(
                    "ERROR",
                    "output",
                    "DB_DO_NOT_FOUND",
                    f"数据库 dispatch、iwt 和 SC fallback 都未查到 DO# {do_no}。",
                    row=raw.source_row if raw else "",
                    pro=raw.pro if raw else "",
                    do_no=do_no,
                    raw_value=do_no,
        )

        calc_rows = dispatch_rows + iwt_rows + sc_rows
        apply_raw_mappings(calc_rows, raw_rows)
        apply_handling_office(calc_rows, dispatch_meta)
        progress(f"[5/8] Querying SC job/customer info for {len(calc_rows)} calculation row(s)...")
        enrich_from_sc_jobinfo(calc_rows, sc_config, validation)
        progress("[6/8] Querying IWT destination/to-firm info...")
        iwt_to = query_iwt_to(
            xc_config,
            unique_preserve(row.do_no for row in calc_rows if row.do_no.upper().startswith("IWT")),
        )

    progress("[7/8] Allocating cost by chargeable weight and building sheets...")
    apply_raw_mappings(calc_rows, raw_rows)
    apply_handling_office(calc_rows, dispatch_meta)
    allocate_costs(calc_rows, validation)
    compare_pro_totals(raw_rows, calc_rows, validation)
    progress(f"[8/8] Saving workbook: {output_path}")
    write_workbook(output_path, raw_rows, calc_rows, dispatch_meta, carrier_lookup, iwt_to, validation)
    return output_path, validation.counts()


def _save_upload(uploaded, label):
    if not uploaded or not uploaded.filename:
        return None
    if not uploaded.filename.lower().endswith(".xlsx") or uploaded.filename.startswith("~$"):
        raise ConverterError("{0} must be a normal .xlsx workbook.".format(label))
    content = uploaded.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise ConverterError("{0} exceeds the 25 MB limit.".format(label))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(uploaded.filename).name)
    path = OUTPUT_DIR / (uuid.uuid4().hex + "_" + name)
    path.write_bytes(content)
    return path


def generate_payload(raw_upload, carrier_upload=None):
    input_path = _save_upload(raw_upload, "Raw workbook")
    if input_path is None:
        raise ConverterError("Select a Raw workbook first.")
    carrier_path = _save_upload(carrier_upload, "Carrier lookup workbook")
    token = uuid.uuid4().hex[:10]
    output_path = OUTPUT_DIR / "{0}_{1}_allocated.xlsx".format(input_path.stem, token)
    try:
        final_path, counts = build_workbook(input_path, output_path, None, carrier_path)
        check = load_workbook(final_path, read_only=True, data_only=False)
        missing = [name for name in ("output", "calculation", "DO", "IWT", "TMS-billing", "Validation_report", "DO_upload", "DO_ARupload") if name not in check.sheetnames]
        check.close()
        if missing:
            raise ConverterError("Generated workbook is missing sheet(s): {0}".format(", ".join(missing)))
    finally:
        input_path.unlink(missing_ok=True)
        if carrier_path:
            carrier_path.unlink(missing_ok=True)
    return {
        "ok": True,
        "filename": final_path.name,
        "download_url": "/download/ord-do-allocation/{0}".format(final_path.name),
        "validation": {key: counts.get(key, 0) for key in ("ERROR", "WARNING", "INFO")},
    }


def output_path_for(filename):
    safe = os.path.basename(str(filename or ""))
    path = (OUTPUT_DIR / safe).resolve()
    return path if safe == filename and path.parent == OUTPUT_DIR.resolve() else None
